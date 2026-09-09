from django.views.decorators.csrf import csrf_exempt
from diabaApp import models
from rest_framework.renderers import JSONRenderer
from rest_framework.parsers import JSONParser
from django.http import HttpResponse
import io
import math, random , calendar
from datetime import datetime, timedelta
from django.conf import settings
from django.db.models import Sum, Count, F, Q, ExpressionWrapper,  OuterRef, Subquery
from django.db.models.functions import Cast, TruncMonth, Coalesce, ExtractMonth, Round, TruncDate
from django.db.models.fields import FloatField
from firebase_admin import messaging
import requests
from django.template.loader import get_template
from django.core.mail import send_mail
from django.http import JsonResponse
from django.db import connection, transaction
from django.db.models.functions import Replace, Lower, Greatest
import uuid
import json
from collections import Counter
import numpy as np
import re
from django.core.cache import cache

import openpyxl, requests
from django.utils import timezone
from django.db.models import Max, Min
from googletrans import Translator
from django.core.files.base import ContentFile, File
from PIL import Image
from io import BytesIO
from django.contrib.postgres.search import TrigramSimilarity
import os 
from django.db.models import Q, F, Count, Case, Sum, When, IntegerField, Value, OuterRef, Subquery, ExpressionWrapper, CharField, FloatField
from django.contrib.postgres.search import TrigramWordSimilarity

from diabaApp.faiss_store import get_faiss_index, preload_clip_categories
from diabaApp.serializer import ProductDetailSerializer, ProductImageSerializer, ProductOtherSpecificationSerializer, ProductModelSerializer, \
    ProductModelVariantSerializerAdmin, VendorProductPriceSerializer, ProductModelVariantSerializer, ProductReviewSerializer, \
    PaymentCargoSliderSerializer, ImportantNoteSerializer, ProductDetailAppSerializer, ProductDataSerializer, \
    ProductDetailAppSerializer, ProductDataSerializer, DummyTagProductNameFrenchSerializer, ProductSearchSerializer, DummyTagProductNameEnglishSerializer, \
    ProductTagSearchSerializer, ProductNameSerializer, ProductCountrySerializer, ProductTagSerializer, RecentlyViewProductSerializer, AdminProductDetailSerializer

import torch
import torchvision.models as torchmodels
import torchvision.transforms as transforms
from PIL import Image
import clip

device = "cuda" if torch.cuda.is_available() else "cpu"
model, preprocess = clip.load("ViT-B/32", device=device)

# Extract vector from uploaded file -> convert file to vector
def extract_features_from_file(file):
    image = Image.open(file).convert("RGB")
    image = preprocess(image).unsqueeze(0).to(device)
    with torch.no_grad():
        features = model.encode_image(image)
    features = features / features.norm(dim=-1, keepdim=True)
    return features.cpu().numpy().astype("float32")

# Extract vector from image URL -> convert db image to vector
def extract_features_from_url(url):
    try:
        response = requests.get(url, timeout=10)
        image = Image.open(BytesIO(response.content)).convert("RGB")
        image = preprocess(image).unsqueeze(0).to(device)
        with torch.no_grad():
            features = model.encode_image(image)
        features = features / features.norm(dim=-1, keepdim=True)
        return features.cpu().numpy().astype("float32")
    except Exception as e:
        print(f"Failed to process {url}: {e}")
        return None


@csrf_exempt
def product_model_variant_create(request):
    python_data={}
    for i,j in request.FILES.items():
        python_data.update({i:j})
    
    for i,j in request.POST.items():
        python_data.update({i:j})

    # print("python_data--=-=-model_name_french-=-=-=>",python_data)


    product_name = (python_data.get('product_name')).strip()
    product_name_french = (python_data.get('product_name_french')).strip()
    category = python_data.get('category')
    subcategory = python_data.get('subcategory')
    # super_subcategory = python_data.get('super_subcategory')
    description = python_data.get('description')
    description_french = python_data.get('description_french')
    unit_of_measure = python_data.get('unit_of_measure')
    material = python_data.get('material')
    length = python_data.get('length')
    height = python_data.get('height')
    width = python_data.get('width')
    weight = python_data.get('weight')
    color = python_data.get('color')
    min_order_quantity = python_data.get('min_order_quantity')
    max_order_quantity = python_data.get('max_order_quantity')
    quantity = python_data.get('quantity')
    product_image_1 = python_data.get('product_image_1')
    product_image_2 = python_data.get('product_image_2')
    product_image_3 = python_data.get('product_image_3')
    product_image_4 = python_data.get('product_image_4')
    product_image_5 = python_data.get('product_image_5')
    product_image_6 = python_data.get('product_image_6')
    product_image_7 = python_data.get('product_image_7')
    product_image_8 = python_data.get('product_image_8')
    
    delay_days_air = python_data.get('delay_days_air',None)
    delay_days_ship = python_data.get('delay_days_ship',None)
    delay_days_express = python_data.get('delay_days_express',None)
    product_type = python_data.get('product_type',None)
    product_packaging = python_data.get('product_packaging',None)
    product_packaging_value = python_data.get('product_packaging_value',None)

    if delay_days_air in [None,'','null']:
        delay_days_air = None

    if delay_days_ship in [None,'','null']:
        delay_days_ship = None

    if delay_days_express in [None,'','null']:
        delay_days_express = None
    
    
    if product_image_1 != None and product_image_1 != 'null' and product_image_1 != '':
        img = Image.open(product_image_1)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")

            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_1.name)
            jpg_name = base_name + ".jpg"
            product_image_1 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_2 != None and product_image_2 != 'null' and product_image_2 != '':
        img = Image.open(product_image_2)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_2.name)
            jpg_name = base_name + ".jpg"
            product_image_2 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_3 != None and product_image_3 != 'null' and product_image_3 != '':
        img = Image.open(product_image_3)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_3.name)
            jpg_name = base_name + ".jpg"
            product_image_3 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_4 != None and product_image_4 != 'null' and product_image_4 != '':
        img = Image.open(product_image_4)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_4.name)
            jpg_name = base_name + ".jpg"
            product_image_4 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_5 != None and product_image_5 != 'null' and product_image_5 != '':
        img = Image.open(product_image_5)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_5.name)
            jpg_name = base_name + ".jpg"
            product_image_5 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_6 != None and product_image_6 != 'null' and product_image_6 != '':
        img = Image.open(product_image_6)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_6.name)
            jpg_name = base_name + ".jpg"
            product_image_6 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_7 != None and product_image_7 != 'null' and product_image_7 != '':
        img = Image.open(product_image_7)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_7.name)
            jpg_name = base_name + ".jpg"
            product_image_7 = ContentFile(buffer.read(), name=jpg_name)

    if product_image_8 != None and product_image_8 != 'null' and product_image_8 != '':
        img = Image.open(product_image_8)
        if img.format == "AVIF":        
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
                
            buffer = BytesIO()
            img.save(buffer, format="JPEG", quality=90)
            buffer.seek(0)
            base_name, ext = os.path.splitext(product_image_8.name)
            jpg_name = base_name + ".jpg"
            product_image_8 = ContentFile(buffer.read(), name=jpg_name)


    price = python_data.get('price')
    discount = python_data.get('discount')
    final_price = python_data.get('final_price')
    vendor = python_data.get('vendor')
    product_verification = python_data.get('product_verification')
    # syncWithDimensions = python_data.get('syncWithDimensions')
    carton_length = python_data.get('carton_length')
    carton_width = python_data.get('carton_width')
    carton_height = python_data.get('carton_height')
    carton_weight = python_data.get('carton_weight')
    # syncWithModelOrVariant = python_data.get('syncWithModelOrVariant')
    available_quantity = python_data.get('available_quantity')
    quantity = python_data.get('quantity')
    refpro = python_data.get('refpro', None)
    reuser = python_data.get('reuser', None)
    shipping_via = python_data.get('shipping_via', None)
    country_of_origin = python_data.get('country_of_origin', None)
    product_video = python_data.get('product_video', None)
    available_country = python_data.get('available_country', None)
    tag = python_data.get('tag', None)
    
    

    other_specification = json.loads(python_data.get('other_specification'))
    
    
    login_flag = python_data.get('login_flag')
    if login_flag == 'Admin':
        status = "Approved"
        product_verification = 'Approved'
        
    else:
        status = "Pending"
        product_verification = 'Pending'

    digits = "123456789"
    random_number = ""
    for i in range(3) :
        random_number += digits[math.floor(random.random() * 9)]

    category_name = models.CategoryDetail.objects.filter(id = category).values_list('category', flat=True)[0][:3].upper()
    subcategory_name = models.SubCategoryDetail.objects.filter(id = subcategory).values_list('subcategory', flat=True)[0][:3].upper()
    product_name_data = (product_name)[:3].upper()
    product_code =  category_name+'_'+subcategory_name+ '_'+ product_name_data + '_' + random_number
    # print(product_code, 'product_codeproduct_code')
    
    if available_country not in [None,"","null"]:
        try:
            available_country = [int(x) for x in available_country.split(',')]
        except Exception as e:

            available_country = json.loads(available_country)

    if tag not in [None,"","null"]:
        try:
            tag = [int(x) for x in tag.split(',')]
        except Exception as e:

            tag = json.loads(tag)


    product = models.ProductDetail.objects.create(
        product_name = product_name,
        product_name_french = product_name_french,
        product_code = product_code,
        vendor_id = vendor,
        category_id = category,
        subcategory_id = subcategory,
        # super_subcategory_id = super_subcategory,
        description = description,
        description_french = description_french,
        unit_of_measure = unit_of_measure,
        material = material,
        length = length,
        height = height,
        width = width,
        weight = weight,
        country_of_origin_id = country_of_origin,
        carton_length = length,
        carton_width = width,
        carton_height = height,
        carton_weight = weight,
        # syncWithModelOrVariant = syncWithModelOrVariant,
        available_quantity = available_quantity,
        quantity = quantity,
        color = color,
        min_order_quantity = min_order_quantity,
        max_order_quantity = max_order_quantity,
        product_image_1 = product_image_1,
        product_image_2 = product_image_2,
        product_image_3 = product_image_3,
        product_image_4 = product_image_4,
        product_image_5 = product_image_5,
        product_image_6 = product_image_6,
        product_image_7 = product_image_7,
        product_image_8 = product_image_8,
        product_video = product_video,
        # price = price,
        discount = discount,
        final_price = final_price,
        status = 'Inactive',
        product_verification = product_verification,
        refpro = refpro,
        reuser = reuser,
        shipping_via = shipping_via,
        delay_days_air = delay_days_air,
        delay_days_ship = delay_days_ship,
        delay_days_express = delay_days_express,
        product_type_id = product_type,
        product_packaging_id = product_packaging,
        product_packaging_value = product_packaging_value,
        
    )
    product.save()
    try: 
        product.available_country.set(available_country)
    except:
        product.available_country.clear()

    try: 
        product.tag.set(tag)
    except:
        product.tag.clear()


    product_id = product.id 


    get_product_image_1 = product.product_image_1
    get_product_image_2 = product.product_image_2
    get_product_image_3 = product.product_image_3
    get_product_image_4 = product.product_image_4
    get_product_image_5 = product.product_image_5
    get_product_image_6 = product.product_image_6
    get_product_image_7 = product.product_image_7
    get_product_image_8 = product.product_image_8

    domain_path = ""


    if get_product_image_1 != None and get_product_image_1 != 'null' and get_product_image_1 != '':
        feat = extract_features_from_url(domain_path+get_product_image_1.url)
        try:
            if feat.size > 0:
                product.product_image_1_vector = feat.tolist()
                product.save()
        except:
            pass
    
    if get_product_image_2 != None and get_product_image_2 != 'null' and get_product_image_2 != '':
        feat = extract_features_from_url(domain_path+get_product_image_2.url)
        try:
            if feat.size > 0:
                product.product_image_2_vector = feat.tolist()
                product.save()
        except:
            pass

    if get_product_image_3 != None and get_product_image_3 != 'null' and get_product_image_3 != '':
        feat = extract_features_from_url(domain_path+get_product_image_3.url)
        try:
            if feat.size > 0:
                product.product_image_3_vector = feat.tolist()
                product.save()
        except:
            pass

    if get_product_image_4 != None and get_product_image_4 != 'null' and get_product_image_4 != '':
        feat = extract_features_from_url(domain_path+get_product_image_4.url)
        try:
            if feat.size > 0:
                product.product_image_4_vector = feat.tolist()
                product.save()
        
        except:
            pass


    if get_product_image_5 != None and get_product_image_5 != 'null' and get_product_image_5 != '':
        feat = extract_features_from_url(domain_path+get_product_image_5.url)
        try:
            if feat.size > 0:
                product.product_image_5_vector = feat.tolist()
                product.save()
        except:
            pass

    if get_product_image_6 != None and get_product_image_6 != 'null' and get_product_image_6 != '':
        feat = extract_features_from_url(domain_path+get_product_image_6.url)
        try:
            if feat.size > 0:
                product.product_image_6_vector = feat.tolist()
                product.save()
        except:
            pass

    if get_product_image_7 != None and get_product_image_7 != 'null' and get_product_image_7 != '':
        feat = extract_features_from_url(domain_path+get_product_image_7.url)
        try:
            if feat.size > 0:
                product.product_image_7_vector = feat.tolist()
                product.save()
        except:
            pass

    if get_product_image_8 != None and get_product_image_8 != 'null' and get_product_image_8 != '':
        feat = extract_features_from_url(domain_path+get_product_image_8.url)
        try:
            if feat.size > 0:
                product.product_image_8_vector = feat.tolist()
                product.save()
        except:
            pass


    for specification in other_specification:
        key = specification['key']
        key_french = specification['key_french']
        value = specification['value']
        value_french = specification['value_french']

        create_specification = models.ProductOtherSpecification.objects.create(
            product_id = product_id,
            key = key,
            key_french = key_french,
            value = value,
            value_french = value_french,
        ).save()


    for i in range(500):
        model_name = f'models[{i}][name]'
        model_name_french = f'models[{i}][model_name_french]'
        if python_data.get(model_name):
            model_name = (python_data.get(model_name)).strip()
            model_name_french = (python_data.get(model_name_french)).strip()
            model_image = python_data.get(f'models[{i}][image]')
            
            if model_image != None and model_image != 'null' and model_image != '':
                img = Image.open(model_image)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(model_image.name)
                    jpg_name = base_name + ".jpg"
                    model_image = ContentFile(buffer.read(), name=jpg_name)

            create_model = models.ProductModel.objects.create(
                product_id = product_id,
                model_name = model_name,
                model_name_french = model_name_french,
                model_image = model_image,
                status = 'Active'
            )
            create_model.save()
            list_price = []
            for j in range(500):
                variant_name = f'models[{i}][variants][{j}][name]'
                variant_name_french = f'models[{i}][variants][{j}][variant_name_french]'
                if python_data.get(variant_name):
                    variant_name = (python_data.get(variant_name)).strip()
                    variant_name_french = (python_data.get(variant_name_french)).strip()
                    variant_image = python_data.get(f'models[{i}][variants][{j}][image]')
                    variant_price = python_data.get(f'models[{i}][variants][{j}][price]')
                    list_price.append(variant_price)

                    if variant_image != None and model_image != 'null' and variant_image != '':
                        img = Image.open(variant_image)
                        if img.format == "AVIF":        
                            if img.mode in ("RGBA", "LA"):
                                img = img.convert("RGB")
                                
                            buffer = BytesIO()
                            img.save(buffer, format="JPEG", quality=90)
                            buffer.seek(0)
                            base_name, ext = os.path.splitext(variant_image.name)
                            jpg_name = base_name + ".jpg"
                            variant_image = ContentFile(buffer.read(), name=jpg_name)

                    create_variant = models.ProductModelVariant.objects.create(
                        product_id = product_id,
                        model_id = create_model.id,
                        name = variant_name,
                        name_french = variant_name_french,
                        image = variant_image,
                        price = variant_price,
                        variant_verification = 'Approved',
                        status = 'Active'

                    )
                    create_variant.save()
                    variant_id = create_variant.id

                    get_vendor_data = f'models[{i}][variants][{j}][vendor]'
                    get_vendor_price = f'models[{i}][variants][{j}][vendor_price]'
                    vendor = python_data.get(get_vendor_data)
                    vendor_price = python_data.get(get_vendor_price)
                   
                    price = models.VendorProductPrice.objects.create(
                        product_id = product_id,
                        model_id = create_model.id,
                        vendor_id = vendor,
                        variant_id = variant_id,
                        price = vendor_price,
                        # quantity = quantity,
                        status = 'Active'
                    ).save()


             
                else:
                    break
        else:
            break    


    smallest_price = min(list_price)
    product_update = models.ProductDetail.objects.get(id = product_id)
    product_update.price = smallest_price
    product_update.save()

    res = {
        'message':'Project create successfully.'        
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)




@csrf_exempt
def product_model_variant_update(request):
    if request.method == 'POST':
        python_data={}
        for i,j in request.FILES.items():
            python_data.update({i:j})
        
        for i,j in request.POST.items():
            python_data.update({i:j})

        id = python_data.get('id')
        product_name = (python_data.get('product_name')).strip()
        product_name_french = (python_data.get('product_name_french')).strip()
        description_french = python_data.get('description_french')
        product_code = python_data.get('product_code', None)
        category = python_data.get('category')
        subcategory = python_data.get('subcategory')
        # super_subcategory = python_data.get('super_subcategory')
        description = python_data.get('description')
        unit_of_measure = python_data.get('unit_of_measure')
        material = python_data.get('material')
        length = python_data.get('length')
        height = python_data.get('height')
        width = python_data.get('width')
        weight = python_data.get('weight')
        color = python_data.get('color')
        min_order_quantity = python_data.get('min_order_quantity')
        max_order_quantity = python_data.get('max_order_quantity')
        quantity = python_data.get('quantity')
        product_image_1 = python_data.get('product_image_1')
        product_image_2 = python_data.get('product_image_2')
        product_image_3 = python_data.get('product_image_3')
        product_image_4 = python_data.get('product_image_4')
        product_image_5 = python_data.get('product_image_5')
        product_image_6 = python_data.get('product_image_6')
        product_image_7 = python_data.get('product_image_7')
        product_image_8 = python_data.get('product_image_8')
        product_video = python_data.get('product_video')
        
        delay_days_air = python_data.get('delay_days_air')
        delay_days_ship = python_data.get('delay_days_ship')
        delay_days_express = python_data.get('delay_days_express')

        if delay_days_air in [None,'','null']:
            delay_days_air = None

        if delay_days_ship in [None,'','null']:
            delay_days_ship = None

        if delay_days_express in [None,'','null']:
            delay_days_express = None

        available_country = python_data.get('available_country')
        tag = python_data.get('tag')
        
        if available_country not in [None,"","null"]:
            try:
                available_country = [int(x) for x in available_country.split(',')]
            except Exception as e:

                available_country = json.loads(available_country)

        if tag not in [None,"","null"]:
            try:
                tag = [int(x) for x in tag.split(',')]
            except Exception as e:

                tag = json.loads(tag)

        if product_image_1 != None and product_image_1 != 'null' and product_image_1 != '':
            try:
                img = Image.open(product_image_1)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")

                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_1.name)
                    jpg_name = base_name + ".jpg"
                    product_image_1 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_2 != None and product_image_2 != 'null' and product_image_2 != '':
            try:
                img = Image.open(product_image_2)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_2.name)
                    jpg_name = base_name + ".jpg"
                    product_image_2 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_3 != None and product_image_3 != 'null' and product_image_3 != '':
            try:
                img = Image.open(product_image_3)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_3.name)
                    jpg_name = base_name + ".jpg"
                    product_image_3 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_4 != None and product_image_4 != 'null' and product_image_4 != '':
            try:
                img = Image.open(product_image_4)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_4.name)
                    jpg_name = base_name + ".jpg"
                    product_image_4 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_5 != None and product_image_5 != 'null' and product_image_5 != '':
            try:
                img = Image.open(product_image_5)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_5.name)
                    jpg_name = base_name + ".jpg"
                    product_image_5 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_6 != None and product_image_6 != 'null' and product_image_6 != '':
            try:
                img = Image.open(product_image_6)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_6.name)
                    jpg_name = base_name + ".jpg"
                    product_image_6 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_7 != None and product_image_7 != 'null' and product_image_7 != '':
            try:
                img = Image.open(product_image_7)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_7.name)
                    jpg_name = base_name + ".jpg"
                    product_image_7 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass
            
        if product_image_8 != None and product_image_8 != 'null' and product_image_8 != '':
            try:
                img = Image.open(product_image_8)
                if img.format == "AVIF":        
                    if img.mode in ("RGBA", "LA"):
                        img = img.convert("RGB")
                        
                    buffer = BytesIO()
                    img.save(buffer, format="JPEG", quality=90)
                    buffer.seek(0)
                    base_name, ext = os.path.splitext(product_image_8.name)
                    jpg_name = base_name + ".jpg"
                    product_image_8 = ContentFile(buffer.read(), name=jpg_name)
            except:
                pass

 
        country_of_origin = python_data.get('country_of_origin')
        
        price = python_data.get('price')
        discount = python_data.get('discount')
        final_price = python_data.get('final_price')
        vendor = python_data.get('vendor')
        product_verification = python_data.get('product_verification')
        refpro = python_data.get('refpro')
        reuser = python_data.get('reuser')

        carton_length = python_data.get('carton_length')
        carton_width = python_data.get('carton_width')
        carton_height = python_data.get('carton_height')
        carton_weight = python_data.get('carton_weight')
        shipping_via = python_data.get('shipping_via')
        status = python_data.get('status')
        product_type = python_data.get('product_type')
        product_packaging = python_data.get('product_packaging')
        product_packaging_value = python_data.get('product_packaging_value')
        
        
        
        
        other_specification = json.loads(python_data.get('other_specification'))

        if vendor in [None,'','null']:
            vendor = None
        # if super_subcategory in [None,'','null']:
        #     super_subcategory = None

        if models.ProductDetail.objects.filter(id = id).exists():
            get_product = models.ProductDetail.objects.get(id = id)
            get_product.product_name = product_name
            get_product.product_name_french = product_name_french
            
            get_product.product_code = product_code
            # get_product.vendor_id = vendor
            get_product.category_id = category
            get_product.subcategory_id = subcategory
            # get_product.super_subcategory_id = super_subcategory
            get_product.description = description
            get_product.description_french = description_french
            get_product.unit_of_measure = unit_of_measure
            get_product.material = material
            get_product.length = length
            get_product.height = height
            get_product.width = width
            get_product.weight = weight
            get_product.color = color
            get_product.min_order_quantity = min_order_quantity
            get_product.max_order_quantity = max_order_quantity
            get_product.quantity = quantity
            get_product.refpro = refpro
            get_product.reuser = reuser
            get_product.country_of_origin_id = country_of_origin

            get_product.carton_length = length
            get_product.carton_width = width
            get_product.carton_height = height
            get_product.carton_weight = weight
            
            if not isinstance(product_image_1, str):
                get_product.product_image_1 = product_image_1
            if not isinstance(product_image_2, str):
                get_product.product_image_2 = product_image_2
            if not isinstance(product_image_3, str):
                get_product.product_image_3 = product_image_3
            if not isinstance(product_image_4, str):
                get_product.product_image_4 = product_image_4
            if not isinstance(product_image_5, str):
                get_product.product_image_5 = product_image_5
            if not isinstance(product_image_6, str):
                get_product.product_image_6 = product_image_6
            if not isinstance(product_image_7, str):
                get_product.product_image_7 = product_image_7
            if not isinstance(product_image_8, str):
                get_product.product_image_8 = product_image_8
            if not isinstance(product_video, str):
                get_product.product_video = product_video


            get_product.price = price
            get_product.discount = discount
            get_product.final_price = final_price
            get_product.product_verification = product_verification
            get_product.shipping_via = shipping_via

            get_product.delay_days_air = delay_days_air
            get_product.delay_days_ship = delay_days_ship
            get_product.delay_days_express = delay_days_express
            get_product.product_type_id = product_type
            get_product.product_packaging_id = product_packaging
            get_product.product_packaging_value = product_packaging_value
            get_product.status = python_data.get('status', get_product.status)


            

            get_product.save()
            try: 
                get_product.available_country.set(available_country)
            except:
                get_product.available_country.clear()
            
            try: 
                get_product.tag.set(tag)
            except:
                get_product.tag.clear()
            
            get_product_image_1 = get_product.product_image_1
            get_product_image_2 = get_product.product_image_2
            get_product_image_3 = get_product.product_image_3
            get_product_image_4 = get_product.product_image_4
            get_product_image_5 = get_product.product_image_5
            get_product_image_6 = get_product.product_image_6
            get_product_image_7 = get_product.product_image_7
            get_product_image_8 = get_product.product_image_8

            domain_path = ""


            if get_product_image_1 != None and get_product_image_1 != 'null' and get_product_image_1 != '':
                feat = extract_features_from_url(domain_path+get_product_image_1.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_1_vector = feat.tolist()
                        get_product.save()
                except:
                    pass
            
            if get_product_image_2 != None and get_product_image_2 != 'null' and get_product_image_2 != '':
                feat = extract_features_from_url(domain_path+get_product_image_2.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_2_vector = feat.tolist()
                        get_product.save()
                except:
                    pass

            if get_product_image_3 != None and get_product_image_3 != 'null' and get_product_image_3 != '':
                feat = extract_features_from_url(domain_path+get_product_image_3.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_3_vector = feat.tolist()
                        get_product.save()
                except:
                    pass

            if get_product_image_4 != None and get_product_image_4 != 'null' and get_product_image_4 != '':
                feat = extract_features_from_url(domain_path+get_product_image_4.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_4_vector = feat.tolist()
                        get_product.save()
                
                except:
                    pass


            if get_product_image_5 != None and get_product_image_5 != 'null' and get_product_image_5 != '':
                feat = extract_features_from_url(domain_path+get_product_image_5.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_5_vector = feat.tolist()
                        get_product.save()
                except:
                    pass

            if get_product_image_6 != None and get_product_image_6 != 'null' and get_product_image_6 != '':
                feat = extract_features_from_url(domain_path+get_product_image_6.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_6_vector = feat.tolist()
                        get_product.save()
                except:
                    pass

            if get_product_image_7 != None and get_product_image_7 != 'null' and get_product_image_7 != '':
                feat = extract_features_from_url(domain_path+get_product_image_7.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_7_vector = feat.tolist()
                        get_product.save()
                except:
                    pass

            if get_product_image_8 != None and get_product_image_8 != 'null' and get_product_image_8 != '':
                feat = extract_features_from_url(domain_path+get_product_image_8.url)
                try:
                    if feat.size > 0:
                        get_product.product_image_8_vector = feat.tolist()
                        get_product.save()
                except:
                    pass


            get_ids = []
            for specification in other_specification:
                key = specification['key']
                key_french = specification['key_french']
                value = specification['value']
                value_french = specification['value_french']
                if specification.get('id'):
                    id = specification['id']
                
                    if models.ProductOtherSpecification.objects.filter(id = id).exists() and id not in [None,'','null']:
                        update_data = models.ProductOtherSpecification.objects.get(id = id)
                        update_data.key = key
                        update_data.key_french = key_french
                        update_data.value = value
                        update_data.value_french = value_french
                        update_data.save()
                        get_ids.append(update_data.id)
                    else:
                        pass
                else:
                    create_specification = models.ProductOtherSpecification.objects.create(
                        product_id = get_product.id,
                        key = key,
                        value = value,
                    )
                    create_specification.save()
                    get_ids.append(create_specification.id)
                                
            delete_specification = models.ProductOtherSpecification.objects.filter(product = get_product.id).exclude(id__in = get_ids)
            delete_specification.delete()

            not_deleted_models_id = []
            for i in range(500):
                model_name = f'models[{i}][name]'
                model_name_french = f'models[{i}][model_name_french]'


                if python_data.get(model_name):
                    model_name = (python_data.get(model_name)).strip()
                    model_name_french = (python_data.get(model_name_french)).strip()
                    model_image = python_data.get(f'models[{i}][image]')
                    model_id = python_data.get(f'models[{i}][id]',None)


                    if models.ProductModel.objects.filter(id = model_id).exists() and model_id not in [None,'','null']:
                        update_model = models.ProductModel.objects.get(id = model_id)
                        update_model.model_name = model_name
                        update_model.model_name_french = model_name_french
                        if not isinstance(model_image, str):
                            # print("MODEL IS HERE")
                            update_model.model_image = model_image

                        update_model.save()
                        not_deleted_models_id.append(update_model.id)

                    else:
                        if model_image != None and model_image != 'null' and model_image != '':
                            img = Image.open(model_image)
                            if img.format == "AVIF":        
                                if img.mode in ("RGBA", "LA"):
                                    img = img.convert("RGB")
                                    
                                buffer = BytesIO()
                                img.save(buffer, format="JPEG", quality=90)
                                buffer.seek(0)
                                base_name, ext = os.path.splitext(model_image.name)
                                jpg_name = base_name + ".jpg"
                                model_image = ContentFile(buffer.read(), name=jpg_name)
                        create_model = models.ProductModel.objects.create(
                            product_id = get_product.id,
                            model_name = model_name,
                            model_name_french = model_name_french,
                            model_image = model_image,
                            status = 'Active'       
                        )
                        create_model.save()
                        model_id = create_model.id
                        not_deleted_models_id.append(model_id)

                    not_deleted_variant_id = []
                    list_price = []
                    # print(True)
                    for j in range(500):
                        variant_name = f'models[{i}][variants][{j}][name]'
                        variant_name_french = f'models[{i}][variants][{j}][variant_name_french]'


                        if python_data.get(variant_name):
                            variant_name = (python_data.get(variant_name)).strip()
                            variant_name_french = (python_data.get(variant_name_french)).strip()
                            variant_image = python_data.get(f'models[{i}][variants][{j}][image]')
                            variant_price = python_data.get(f'models[{i}][variants][{j}][price]')
                            variant_id = python_data.get(f'models[{i}][variants][{j}][id]',None)
                            list_price.append(variant_price)

                            

                            if variant_id not in [None,'','null'] and models.ProductModelVariant.objects.filter(id = variant_id).exists():
                                update_variant = models.ProductModelVariant.objects.get(id = variant_id)
                                update_variant.name = variant_name
                                update_variant.name_french = variant_name_french
                                
                                update_variant.price = variant_price

                                if not isinstance(variant_image, str):
                                    # print("VARIANT IS HERE")
                                    update_variant.image = variant_image
                                
                                update_variant.save()
                                not_deleted_variant_id.append(update_variant.id)

                            else:
                                if variant_image != None and variant_image != 'null' and variant_image != '':
                                    img = Image.open(variant_image)
                                    if img.format == "AVIF":        
                                        if img.mode in ("RGBA", "LA"):
                                            img = img.convert("RGB")
                                            
                                        buffer = BytesIO()
                                        img.save(buffer, format="JPEG", quality=90)
                                        buffer.seek(0)
                                        base_name, ext = os.path.splitext(variant_image.name)
                                        jpg_name = base_name + ".jpg"
                                        variant_image = ContentFile(buffer.read(), name=jpg_name)


                                create_variant = models.ProductModelVariant.objects.create(
                                    product_id = get_product.id,
                                    model_id = model_id,
                                    name = variant_name,
                                    name_french = variant_name_french,
                                    
                                    image = variant_image,
                                    price = variant_price,
                                    variant_verification = 'Approved'

                                )
                                create_variant.save()
                                
                                not_deleted_variant_id.append(create_variant.id)
                        
                        else:
                            break
                    
                    delete_variant = models.ProductModelVariant.objects.filter(model = model_id).exclude(id__in = not_deleted_variant_id)
                    delete_variant.delete()

                else:
                    break
            
            delete_model = models.ProductModel.objects.filter(product_id = get_product.id).exclude(id__in = not_deleted_models_id)
            delete_model.delete()

            # print(list_price)
            smallest_price = min(list_price)
            # print(smallest_price, 'smallest_pricesmallest_pricesmallest_price')
            get_product = models.ProductDetail.objects.get(id = id)
            get_product.price = smallest_price
            get_product.save()

            res={
                'message':'Product Updated Successfully'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)


@csrf_exempt
def product_model_status_update(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)
    id = python_data.get('id')

    product = models.ProductModel.objects.get(id = id)
    product.status = python_data.get('status', product.status)
    product.save()

    res = {
        'message':'Product verification update successfully.'        
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)

    
@csrf_exempt
def product_variant_status_update(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)

    # print("python_Data-=-=-->",python_data)
    id = python_data.get('id')

    product = models.ProductModelVariant.objects.get(id = id)
    product.status = python_data.get('status', product.status)
    product.save()

    res = {
        'message':'Product verification update successfully.'        
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)


@csrf_exempt
def product_vendor_status_update(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)
    id = python_data.get('id')

    product = models.VendorProductPrice.objects.get(id = id)
    product.status = python_data.get('status', product.status)
    product.save()

    res = {
        'message':'Product verification update successfully.'        
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)



@csrf_exempt
def product_detail_admin(request):
    if request.method == 'POST':
        python_data = JSONParser().parse(io.BytesIO(request.body))
        id = python_data.get('id')
        
        product_check = models.ProductDetail.objects.filter(id = id).count()
        if product_check == 1:
            product = models.ProductDetail.objects.get(id = id)
            data = ProductDetailSerializer(product).data

            get_image = models.ProductImage.objects.filter(product = id)
            image_serializer = ProductImageSerializer(get_image, many=True).data

            get_other_specification = models.ProductOtherSpecification.objects.filter(product = id)
            other_serializer = ProductOtherSpecificationSerializer(get_other_specification, many=True).data

            list_model = []

            all_model = models.ProductModel.objects.filter(product = id).order_by('-id')
            for model in all_model:
                model = models.ProductModel.objects.get(id = model.id)
                model_serializer = ProductModelSerializer(model).data
                get_variants = models.ProductModelVariant.objects.filter(model = model.id).order_by('-id')
                list_variant = []
                for variant in get_variants:
                    variants_data = models.ProductModelVariant.objects.get(id = variant.id)
                    variant_serilizer = ProductModelVariantSerializerAdmin(variants_data).data

                    vendor_data = models.VendorProductPrice.objects.filter(variant= variant.id).order_by('-id')
                    price_serializer = VendorProductPriceSerializer(vendor_data, many=True).data
                    
                    variant_serilizer.update({'vendor_list':price_serializer})
                    list_variant.append(variant_serilizer)

                model_serializer.update({'variants':list_variant})    
                
                list_model.append(model_serializer)
            
            data.update({'image':image_serializer, 'other_data':other_serializer})

        res={
            'data':data,
            'models':list_model
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)



@csrf_exempt
def product_detail_vendor(request):
    if request.method == 'POST':
        python_data = JSONParser().parse(io.BytesIO(request.body))
        id = python_data.get('id')
        vendor = python_data.get('vendor')
        
        product_check = models.ProductDetail.objects.filter(id = id).count()
        if product_check == 1:
            product = models.ProductDetail.objects.get(id = id)
            data = ProductDetailSerializer(product).data

            get_image = models.ProductImage.objects.filter(product = id)
            image_serializer = ProductImageSerializer(get_image, many=True).data

            get_other_specification = models.ProductOtherSpecification.objects.filter(product = id)
            other_serializer = ProductOtherSpecificationSerializer(get_other_specification, many=True).data

            list_model = []

            all_model = models.ProductModel.objects.filter(product = id).order_by('-id')
            for model in all_model:
                model = models.ProductModel.objects.get(id = model.id)
                model_serializer = ProductModelSerializer(model).data
                get_variants = models.ProductModelVariant.objects.filter(model = model.id).order_by('-id')
                list_variant = []
                for variant in get_variants:
                    variants_data = models.ProductModelVariant.objects.get(id = variant.id)
                    variant_serilizer = ProductModelVariantSerializer(variants_data).data

                    # print(variant.id,'<-----variant.id, vendor---->', vendor)

                    vendor_data = models.VendorProductPrice.objects.get(variant= variant.id, vendor = vendor)
                    
                    variant_serilizer.update({'vendor_price':vendor_data.price})
                    list_variant.append(variant_serilizer)

                model_serializer.update({'variants':list_variant})    
                
                list_model.append(model_serializer)
            
            data.update({'image':image_serializer, 'other_data':other_serializer})

        res={
            'data':data,
            'models':list_model
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)



@csrf_exempt
def product_detail_app(request):
    if request.method == 'POST':
        python_data = JSONParser().parse(io.BytesIO(request.body))
        print(python_data, 'python_data customer ')
        id = python_data.get('id')
        customer = python_data.get('customer')

        currencyCode = python_data.get('currencyCode',None)
        country = python_data.get('countryName','Egypt')
        ip_address = python_data.get('ip_address')
        
        price_update = models.DailyPrice.objects.last()
        # daily_price = DailyPriceSerializer(price_update).data

        get_delivery_data = models.DeliveryDayDetail.objects.first()
    
        now = datetime.now()
        date_time = now.strftime('%Y-%m-%d, %H:%M:%S')

        product = models.ProductDetail.objects.get(id = id)
        
        delay_days_air = 0
        delay_days_ship = 0
        delay_days_express = 0
        if product.delay_days_air not in [None,'','null']:
            delay_days_air = product.delay_days_air
        if product.delay_days_ship not in [None,'','null']:
            delay_days_ship = product.delay_days_ship
        if product.delay_days_express not in [None,'','null']:
            delay_days_express = product.delay_days_express

     

        express_date_now = now + timedelta(days=get_delivery_data.express_delivery + delay_days_express )
        express_date = express_date_now.strftime('%d').lower()
        express_month = express_date_now.strftime('%B').lower()
        
        air_date_now = now + timedelta(days=get_delivery_data.air_delivery + delay_days_air)
        air_date = air_date_now.strftime('%d').lower()
        air_month = air_date_now.strftime('%B').lower()
        
        ship_date_now = now + timedelta(days=get_delivery_data.ship_delivery + delay_days_ship)
        ship_date = ship_date_now.strftime('%d').lower()
        ship_month = ship_date_now.strftime('%B').lower()

        month_map = {
            "january": "janvier",
            "february": "février",
            "march": "mars",
            "april": "avril",
            "may": "mai",
            "june": "juin",
            "july": "juillet",
            "august": "août",
            "september": "septembre",
            "october": "octobre",
            "november": "novembre",
            "december": "décembre"
        }
        express_fr_month = month_map.get(express_month)
        air_fr_month = month_map.get(air_month)
        ship_fr_month = month_map.get(ship_month)
        
        express_en = f"before {express_date} {express_month}"
        express_fr = f"avant le {express_date} {express_fr_month}"
        air_en = f"before {air_date} {air_month}"
        air_fr = f"avant le {air_date} {air_fr_month}"
        ship_en = f"before {ship_date} {ship_month}"
        ship_fr = f"avant le {ship_date} {ship_fr_month}"




        try:
            if currencyCode not in [None,'','null']:
                if customer:
                    country = models.CountryWithCurrency.objects.filter(currency_code = currencyCode).first().country_name
                    
                    user = models.CustomerDetail.objects.get(id = customer)
                    delivery_country = user.country
                else:
                    delivery_country = country
                    country = models.CountryWithCurrency.objects.filter(currency_code = currencyCode).first().country_name

            elif customer:
                user = models.CustomerDetail.objects.get(id = customer)
                countryCode = user.countryCode

                if models.CountryWithCurrency.objects.filter(country_calling_code = countryCode).exists():
                    country = models.CountryWithCurrency.objects.get(country_calling_code = countryCode).country_name
                    delivery_country = country
                else:
                    if country not in [None,'','null']:
                        delivery_country = country
                    else:
                        try:
                            response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                            response = response.json()
                            country = response.get("country_name")
                            # print("response====>",response)
                            delivery_country = country
                            if country in [None,'','null']:
                                country = "Egypt"

                        except Exception as e:
                            # print("Errroooorrr----->",e)
                            country = "Egypt"
                            delivery_country = country
                        

            else:
                if country in [None,'','null']:
                    try:
                        response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                        response = response.json()
                        country = response.get("country_name")
                        # print("response====>",response)

                        delivery_country = country

                        if country in [None,'','null']:
                            country = "Egypt"

                    except Exception as e:
                        # print("Errroooorrr----->",e)
                        country = "Egypt"
                else:
                    try:
                        response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                        response = response.json()
                        delivery_country = response.get("country_name")
                        print("response====>",response)

                        if delivery_country in [None,'','null']:
                            delivery_country = country

                    except Exception as e:
                        # print("Errroooorrr----->",e)
                        delivery_country = country
                        
                    country = country
        except Exception as e:
            # print("Error-=-=-=--->",e)
            country = 'Egypt'
            delivery_country = country


        # print(delivery_country, 'delivery_countrydelivery_country')
        product_review_data = models.ProductReview.objects.filter(product = id).order_by('-id')[:5]
        review_serialiser = ProductReviewSerializer(product_review_data, many=True).data
        
        payment_data = models.PaymentCargoSlider.objects.filter(status = 'Active')
        slider_serialser = PaymentCargoSliderSerializer(payment_data, many=True).data

        important_note = models.ImportantNote.objects.filter(status = 'Active').first()
        # print("important_note---->",important_note)
        if important_note:
            note_serialser = ImportantNoteSerializer(important_note).data
        else:
            note_serialser=None
        
        # print("note_serialser---->",note_serialser)
        

        total_review = models.ProductReview.objects.filter(product = id).count()
        average_review = 0
        if total_review > 0:
            get_rating = list(models.ProductReview.objects.filter(product = id).values_list('rating', flat=True))
            average_review = round(sum(get_rating) / total_review, 1)
            # print(get_rating, 'get_rating', average_review)   

        if customer:
            product_check = models.ProductDetail.objects.filter(id = id).count()
            if product_check == 1:
                wishlist_count = models.WishlistDetail.objects.filter(customer = customer, product = id).count()
                if wishlist_count == 1:
                    wishlist_flag = True
                else:
                    wishlist_flag = False

                product = models.ProductDetail.objects.get(id = id)
                data = ProductDetailAppSerializer(product, context = {'country':country,'delivery_country':delivery_country}).data                
                subcategory = product.subcategory.id

                get_image = models.ProductImage.objects.filter(product = id)
                image_serializer = ProductImageSerializer(get_image, many=True).data

                get_other_specification = models.ProductOtherSpecification.objects.filter(product = id)
                other_serializer = ProductOtherSpecificationSerializer(get_other_specification, many=True).data

                list_model = []
                total_quantity = 0 
                all_model = models.ProductModel.objects.filter(product = id).order_by('-id')
                for model in all_model:
                    model = models.ProductModel.objects.get(id = model.id)
                    model_serializer = ProductModelSerializer(model).data
                    get_variants = models.ProductModelVariant.objects.filter(model = model.id).order_by('-id')
                    list_variant = []
                    for variant in get_variants:
                        variants_data = models.ProductModelVariant.objects.get(id = variant.id)
                        variant_serilizer = ProductModelVariantSerializer(variants_data, context = {'country':country}).data

                        # check for delay note:

                        # vendor_qs = models.VendorProductPrice.objects.filter(
                        #     variant=variants_data
                        # ).values_list('vendor', flat=True).distinct()

                        # vendor_count = vendor_qs.count()

                        # delay_notes_qs = models.VendorDelayNote.objects.filter(
                        #     vendor__in=vendor_qs,
                        #     start_at__lte=now,
                        #     end_at__gte=now,
                        #     status="Active"
                        # )

                        # delay_note_vendor_count = delay_notes_qs.count()

                        # note = None

                        # if vendor_count > 0 and delay_note_vendor_count == vendor_count:
                        #     note = delay_notes_qs.order_by('-id').values_list('note', flat=True).first()

                        # variant_serilizer.update({'delay_note': note})                        
                        
                        check_varient = models.CartDetail.objects.filter(customer = customer, product = id, variant = variant.id).count()
                        # print(check_varient, 'check_varient')
                        if check_varient != 0:
                            cart_quantity = models.CartDetail.objects.filter(customer = customer, product = id, variant = variant.id).values_list('quantity', flat=True)[0]
                            total_quantity += int(cart_quantity)
                            variant_serilizer.update({'cart_quantity':cart_quantity})
                        else:
                            variant_serilizer.update({'cart_quantity':0})

                        # variant_serilizer.update({'vendor_list':price_serializer})
                        list_variant.append(variant_serilizer)


                    model_serializer.update({'variants':list_variant})    
                    
                    list_model.append(model_serializer)
                
                data.update({'image':image_serializer, 'other_data':other_serializer})
                # print(data, 'afasf')
                
                all_related_product = models.ProductDetail.objects.filter(subcategory = subcategory, status='Active')[:10]
                list_related_product = []
                for related_product in all_related_product:
                    current_id = related_product.id
                    if product.id != current_id:
                        current_data = models.ProductDetail.objects.get(id = current_id)
                        product_serializer = ProductDataSerializer(current_data,context = {'country':country}).data
                        list_related_product.append(product_serializer)

                # print(list_related_product, 'list_related_product')

                check_review_product = models.RecentlyViewProduct.objects.filter(customer = customer, product = id).count()
                if check_review_product == 1:
                    get_review_product = models.RecentlyViewProduct.objects.get(customer = customer, product = id)
                    get_review_product.delete()

                    get_review_product = models.RecentlyViewProduct.objects.create(
                        customer_id = customer, 
                        product_id = id,
                        created_at = date_time

                    ).save()

                elif check_review_product == 0:
                    get_review_product = models.RecentlyViewProduct.objects.create(
                        customer_id = customer, 
                        product_id = id,
                        created_at = date_time

                    ).save()


            # air_fr = f'avant le {}'
                                    
            res={
                'data':data,
                'models':list_model, 
                'wishlist_flag':wishlist_flag,
                # 'cart_detail':cart_detail_serialiser, 
                # 'daily_price':daily_price,
                'list_related_product':list_related_product,
                'express_en':express_en,
                'express_fr':express_fr,
                'air_en':air_en,
                'air_fr':air_fr,
                'ship_en':ship_en,
                'ship_fr':ship_fr,
                'total_review':total_review,
                'average_review':average_review,
                'review_data':review_serialiser,
                'payment_cargo_slider':slider_serialser,
                'important_note':note_serialser,
                'total_quantity':total_quantity
                
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)

        else:
            product_check = models.ProductDetail.objects.filter(id = id).count()
            if product_check == 1:
                
                product = models.ProductDetail.objects.get(id = id)
                data = ProductDetailAppSerializer(product,context = {'country':country,'delivery_country':delivery_country}).data

                # print('data---->',data)

                subcategory = product.subcategory.id

                get_image = models.ProductImage.objects.filter(product = id)
                image_serializer = ProductImageSerializer(get_image, many=True).data

                get_other_specification = models.ProductOtherSpecification.objects.filter(product = id)
                other_serializer = ProductOtherSpecificationSerializer(get_other_specification, many=True).data 
                

                list_model = []

                all_model = models.ProductModel.objects.filter(product = id).order_by('-id')
                for model in all_model:
                    model = models.ProductModel.objects.get(id = model.id)
                    model_serializer = ProductModelSerializer(model).data
                    get_variants = models.ProductModelVariant.objects.filter(model = model.id).order_by('-id')
                    list_variant = []
                    for variant in get_variants:
                        variants_data = models.ProductModelVariant.objects.get(id = variant.id)
                        variant_serilizer = ProductModelVariantSerializer(variants_data,context = {'country':country}).data
                        
                        # variant_serilizer.update({'vendor_list':price_serializer})
                        
                        # check for delay note:

                        # vendor_qs = models.VendorProductPrice.objects.filter(
                        #     variant=variants_data
                        # ).values_list('vendor', flat=True).distinct()

                        # vendor_count = vendor_qs.count()

                        # delay_notes_qs = models.VendorDelayNote.objects.filter(
                        #     vendor__in=vendor_qs,
                        #     start_at__lte=now,
                        #     end_at__gte=now,
                        #     status="Active"
                        # )

                        # delay_note_vendor_count = delay_notes_qs.count()

                        # note = None

                        # if vendor_count > 0 and delay_note_vendor_count == vendor_count:
                        #     note = delay_notes_qs.order_by('-id').values_list('note', flat=True).first()

                        # variant_serilizer.update({'delay_note': note})

                        list_variant.append(variant_serilizer)

                    model_serializer.update({'variants':list_variant})
                    
                    list_model.append(model_serializer)
                
                data.update({'image':image_serializer, 'other_data':other_serializer})

                all_related_product = models.ProductDetail.objects.filter(subcategory = subcategory, status='Active')[:10]
                list_related_product = []
                for related_product in all_related_product:
                    current_id = related_product.id
                    if product.id != current_id:
                        current_data = models.ProductDetail.objects.get(id = current_id)
                        product_serializer = ProductDataSerializer(current_data,context = {'country':country}).data
                        list_related_product.append(product_serializer)

               
            

            res={
                'data':data,
                'models':list_model,
                # 'daily_price':daily_price,
                'list_related_product':list_related_product,
                'express_en':express_en,
                'express_fr':express_fr,
                'air_en':air_en,
                'air_fr':air_fr,
                'ship_en':ship_en,
                'ship_fr':ship_fr,
                'total_review':total_review,
                'average_review':average_review,
                'review_data':review_serialiser,
                'payment_cargo_slider':slider_serialser,
                'important_note':note_serialser
                
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)


@csrf_exempt
def product_search(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))

        # print("python_data--product_search-->",python_data)

        search = python_data.get('search')
        currencyCode = python_data.get('currencyCode',None)
        countryCallingCode = python_data.get('countryCallingCode',None)
        country = python_data.get('countryName','Egypt')
        ip_address = python_data.get('ip_address')
        user_id = python_data.get('user_id')

        try:
            if currencyCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(currency_code = currencyCode).first().country_name
            elif countryCallingCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(country_calling_code = countryCallingCode).first().country_name
                
            elif user_id:
                user = models.CustomerDetail.objects.get(id = user_id)
                countryCode = user.countryCode

                if models.CountryWithCurrency.objects.filter(country_calling_code = countryCode).exists():
                    country = models.CountryWithCurrency.objects.get(country_calling_code = countryCode).country_name
                else:
                    if country not in [None,'','null']:
                        country = country
                    else:
                        try:
                            response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                            response = response.json()
                            country = response.get("country_name")
                            print("response====>",response)

                            if country in [None,'','null']:
                                country = "Egypt"

                        except Exception as e:
                            print("Errroooorrr----->",e)
                            country = "Egypt"
                        

            else:
                if country in [None,'','null']:
                    try:
                        response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                        response = response.json()
                        country = response.get("country_name")
                        print("response====>",response)

                        if country in [None,'','null']:
                            country = "Egypt"

                    except Exception as e:
                        print("Errroooorrr----->",e)
                        country = "Egypt"
                else:
                    country = country
        except Exception as e:
            print("Error-=-=-=--->",e)
            country = 'Egypt'


        
        serach_query = Q()

#         # serach_query &= Q(product_name__icontains = search) | Q(product_name_french__icontains = search) | Q(category__category__icontains = search) | Q(category__category_french__icontains = search)
#         serach_query &= Q(product_name__icontains = search, status='Active') | Q(product_name_french__icontains = search, status='Active') | Q(refpro__icontains = search, status='Active')

# # __istartswith
#         product_list = models.ProductDetail.objects.filter(serach_query).distinct()

        ### code check ## 
        # product_list = models.ProductDetail.objects.filter(Q(product_name__icontains=search) |Q(product_name_french__icontains=search) |Q(refpro__icontains=search),status='Active').annotate(
        # priority=Case(
        # # STARTSWITH (Highest priority)
        # When(
        #     Q(product_name__istartswith=search) |
        #     Q(product_name_french__istartswith=search) |
        #     Q(refpro__istartswith=search),
        #     then=0
        # ),
        # # CONTAINS (Lower priority)
        # default=1,
        # output_field=IntegerField(),
        # )).order_by("priority", "product_name")


        search_clean = search.replace(" ", "").lower()

        product_list = (
            models.ProductDetail.objects
            .annotate(
                clean_name=Lower(
                    Replace('product_name', Value(' '), Value(''))
                ),
            )
            .annotate(
                similarity=TrigramSimilarity('clean_name', search_clean)
            )
            .filter(
                similarity__gt=0.25,
                status='Active'
            )
            .order_by('-similarity')
        )

        product_serializer = ProductDataSerializer(product_list, many=True,context = {'country':country}).data

        res={
            'data':product_serializer
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)



def get_product_result(
    tag,
    country,
    search_by,
    search,
    country_name
):
    search_query = Q()

    search_query &= Q(available_country__id=country)

    # -----------------------------
    # TAG FILTER
    # -----------------------------
    if tag not in [None, '', 'null']:

        if search_by == "Tag":

            search_query &= Q(tag__id=tag)

        elif search_by == "product_english":

            product = models.ProductDetail.objects.filter(
                id=tag
            ).first()

            if product:
                search_query &= Q(
                    product_name__unaccent__icontains=
                    product.product_name
                )

        elif search_by == "product_french":

            product = models.ProductDetail.objects.filter(
                id=tag
            ).first()

            if product:
                search_query &= Q(
                    product_name_french__unaccent__icontains=
                    product.product_name_french
                )

    search_clean = ""

    if search not in [None, '', 'null']:
        search = search.strip()
        search_clean = search.replace(" ", "").lower()

    products = models.ProductDetail.objects.annotate(

        clean_product_name=Lower(
            Replace('product_name', Value(' '), Value(''))
        ),

        clean_product_name_french=Lower(
            Replace('product_name_french', Value(' '), Value(''))
        ),

        clean_refpro=Lower(
            Replace('refpro', Value(' '), Value(''))
        )
    )

    # -----------------------------
    # SEARCH FILTER
    # -----------------------------
    # if search_clean:

    #     products = products.annotate(

    #         similarity_name=TrigramSimilarity(
    #             'product_name',
    #             search
    #         ),

    #         similarity_name_french=TrigramSimilarity(
    #             'product_name_french',
    #             search
    #         ),

    #         similarity_refpro=TrigramSimilarity(
    #             'refpro',
    #             search
    #         )
    #     )

    #     search_words = search.lower().split()

    #     word_query = Q()

    #     for word in search_words:

    #         word_query |= Q(
    #             product_name__unaccent__icontains=word
    #         )

    #         word_query |= Q(
    #             product_name_french__unaccent__icontains=word
    #         )

    #         word_query |= Q(
    #             refpro__unaccent__icontains=word
    #         )

    #     search_query &= (
    #         word_query |
    #         Q(similarity_name__gt=0.15) |
    #         Q(similarity_name_french__gt=0.15) |
    #         Q(similarity_refpro__gt=0.15)
    #     )

    if search_clean:

        import re

        normalized_search = re.sub(
            r'(.)\1+',
            r'\1',
            search.lower()
        )

        products = products.annotate(

            similarity_name=TrigramSimilarity(
                'product_name',
                normalized_search
            ),

            similarity_name_word=TrigramWordSimilarity(
                normalized_search,
                'product_name'
            ),

            similarity_name_french=TrigramSimilarity(
                'product_name_french',
                normalized_search
            ),

            similarity_name_french_word=TrigramWordSimilarity(
                normalized_search,
                'product_name_french'
            ),

            similarity_refpro=TrigramSimilarity(
                'refpro',
                normalized_search
            ),

            similarity_refpro_word=TrigramWordSimilarity(
                normalized_search,
                'refpro'
            )

        ).annotate(

            best_similarity=Greatest(
                'similarity_name',
                'similarity_name_word',
                'similarity_name_french',
                'similarity_name_french_word',
                'similarity_refpro',
                'similarity_refpro_word'
            )
        )

        search_words = normalized_search.split()

        word_query = Q()

        for word in search_words:

            word_query |= Q(
                product_name__unaccent__icontains=word
            )

            word_query |= Q(
                product_name_french__unaccent__icontains=word
            )

            word_query |= Q(
                refpro__unaccent__icontains=word
            )

        search_query &= (
            word_query |
            Q(best_similarity__gt=0.30)
        )

    tag_result = (
        products
        .filter(
            search_query,
            status='Active'
        )
        .distinct()
        .order_by(
            '-best_similarity',
            'product_name'
        )
    )

    tag_result_serializer = ProductSearchSerializer(
        tag_result,
        many=True,
        context={'country': country_name}
    ).data

    # -----------------------------
    # FALLBACK SEARCH
    # -----------------------------
    if not tag_result_serializer:

        try:

            tag_name = (
                models.ProductTag.objects
                .get(id=tag)
                .tag
            )

            fallback_search = (
                tag_name
                .replace(" ", "")
                .lower()
            )

            without_tag_result = (
                products
                .filter(
                    (
                        Q(
                            clean_product_name__unaccent__icontains=
                            fallback_search
                        ) |
                        Q(
                            clean_product_name_french__unaccent__icontains=
                            fallback_search
                        ) |
                        Q(
                            clean_refpro__unaccent__icontains=
                            fallback_search
                        )
                    ),
                    available_country__id=country,
                    status='Active'
                )
                .distinct()
            )

            tag_result_serializer = ProductSearchSerializer(
                without_tag_result,
                many=True,
                context={
                    'country': country_name
                }
            ).data

        except Exception:

            without_tag_result = (
                products
                .filter(
                    (
                        Q(
                            clean_product_name__unaccent__icontains=
                            search_clean
                        ) |
                        Q(
                            clean_product_name_french__unaccent__icontains=
                            search_clean
                        ) |
                        Q(
                            clean_refpro__unaccent__icontains=
                            search_clean
                        )
                    ),
                    available_country__id=country,
                    status='Active'
                )
                .distinct()
            )

            tag_result_serializer = ProductSearchSerializer(
                without_tag_result,
                many=True,
                context={
                    'country': country_name
                }
            ).data

    return tag_result_serializer




@csrf_exempt
def product_search_by_tag(request):

    if request.method != "POST":
        return HttpResponse(
            JSONRenderer().render({
                'message': 'Invalid request method'
            }),
            content_type='application/json',
            status=405
        )

    try:

        python_data = JSONParser().parse(
            io.BytesIO(request.body)
        )

        search = python_data.get('search')
        country = python_data.get('country')
        country_name = python_data.get(
            'countryName',
            'Egypt'
        )

        search = str(search or "").strip()

        if search.lower() == 'null':
            search = ""

        search_clean = (
            search
            .replace(" ", "")
            .replace("-", "")
            .replace("_", "")
            .lower()
        )

        serializer_context = {
            'country': country_name
        }

        # =====================================
        # TAG SEARCH
        # =====================================

        tag_queryset = (
            models.ProductTag.objects
            .annotate(
                clean_tag=Lower(
                    Replace(
                        Replace(
                            Replace(
                                'tag',
                                Value(' '),
                                Value('')
                            ),
                            Value('-'),
                            Value('')
                        ),
                        Value('_'),
                        Value('')
                    )
                ),
                priority=Case(
                    When(
                        clean_tag=search_clean,
                        then=0
                    ),
                    When(
                        clean_tag__startswith=search_clean,
                        then=1
                    ),
                    When(
                        tag__istartswith=search,
                        then=2
                    ),
                    When(
                        clean_tag__icontains=search_clean,
                        then=3
                    ),
                    default=4,
                    output_field=IntegerField()
                )
            )
            .filter(
                Q(tag__unaccent__icontains=search) |
                Q(clean_tag__icontains=search_clean),
                productdetail__available_country__id=country
            )
            .distinct()
            .order_by(
                'priority',
                'clean_tag',
                'tag'
            )
        )

        serializer = ProductTagSearchSerializer(
            tag_queryset,
            many=True,
            context=serializer_context
        ).data

        # =====================================
        # TAG FOUND
        # =====================================

        if serializer:

            products = get_product_result(
                tag=serializer[0]['id'],
                country=country,
                search_by="Tag",
                search=search,
                country_name=country_name
            )

            return HttpResponse(
                JSONRenderer().render({
                    'status': True,
                    'search_by': 'Tag',
                    'data': serializer,
                    'products': products
                }),
                content_type='application/json',
                status=200
            )

        # =====================================
        # ENGLISH PRODUCT SEARCH
        # =====================================

        product_queryset = (
            models.ProductDetail.objects
            .annotate(
                similarity_name=TrigramWordSimilarity(
                    search,
                    'product_name'
                ),

                similarity_ref=TrigramWordSimilarity(
                    search,
                    'refpro'
                ),

                clean_product_name=Lower(
                    Replace(
                        Replace(
                            Replace(
                                'product_name',
                                Value(' '),
                                Value('')
                            ),
                            Value('-'),
                            Value('')
                        ),
                        Value('_'),
                        Value('')
                    )
                ),

                clean_refpro=Lower(
                    Replace(
                        Replace(
                            Replace(
                                'refpro',
                                Value(' '),
                                Value('')
                            ),
                            Value('-'),
                            Value('')
                        ),
                        Value('_'),
                        Value('')
                    )
                ),
            )
            .annotate(
                best_similarity=Greatest(
                    'similarity_name',
                    'similarity_ref'
                )
            )
            .filter(
                available_country__id=country,
                status='Active'
            )
            .filter(
                Q(product_name__unaccent__icontains=search) |
                Q(clean_product_name__icontains=search_clean) |
                Q(refpro__unaccent__icontains=search) |
                Q(clean_refpro__icontains=search_clean) |
                Q(best_similarity__gt=0.10)
            )
            .distinct()
            .order_by('-best_similarity', 'product_name')[:15]
        )

        serializer = DummyTagProductNameEnglishSerializer(
            product_queryset,
            many=True,
            context=serializer_context
        ).data

        if serializer:

            products = get_product_result(
                tag=serializer[0]['id'],
                country=country,
                search_by="product_english",
                search=search,
                country_name=country_name
            )

            return HttpResponse(
                JSONRenderer().render({
                    'status': True,
                    'search_by': 'product_english',
                    'data': serializer,
                    'products': products
                }),
                content_type='application/json',
                status=200
            )

        # =====================================
        # FRENCH PRODUCT SEARCH
        # =====================================

        product_queryset = (
            models.ProductDetail.objects
            .annotate(
                similarity_name=TrigramWordSimilarity(
                    search,
                    'product_name_french'
                ),

                similarity_ref=TrigramWordSimilarity(
                    search,
                    'refpro'
                ),

                clean_product_name_french=Lower(
                    Replace(
                        Replace(
                            Replace(
                                'product_name_french',
                                Value(' '),
                                Value('')
                            ),
                            Value('-'),
                            Value('')
                        ),
                        Value('_'),
                        Value('')
                    )
                ),

                clean_refpro=Lower(
                    Replace(
                        Replace(
                            Replace(
                                'refpro',
                                Value(' '),
                                Value('')
                            ),
                            Value('-'),
                            Value('')
                        ),
                        Value('_'),
                        Value('')
                    )
                ),
            )
            .annotate(
                best_similarity=Greatest(
                    'similarity_name',
                    'similarity_ref'
                )
            )
            .filter(
                available_country__id=country,
                status='Active'
            )
            .filter(
                Q(product_name_french__unaccent__icontains=search) |
                Q(clean_product_name_french__icontains=search_clean) |
                Q(refpro__unaccent__icontains=search) |
                Q(clean_refpro__icontains=search_clean) |
                Q(best_similarity__gt=0.10)
            )
            .distinct()
            .order_by('-best_similarity', 'product_name_french')[:15]
        )

        serializer = DummyTagProductNameFrenchSerializer(
            product_queryset,
            many=True,
            context=serializer_context
        ).data

        if serializer:

            products = get_product_result(
                tag=serializer[0]['id'],
                country=country,
                search_by="product_french",
                search=search,
                country_name=country_name
            )

            return HttpResponse(
                JSONRenderer().render({
                    'status': True,
                    'search_by': 'product_french',
                    'data': serializer,
                    'products': products
                }),
                content_type='application/json',
                status=200
            )

        return HttpResponse(
            JSONRenderer().render({
                'status': True,
                'search_by': '',
                'data': [],
                'products': []
            }),
            content_type='application/json',
            status=200
        )

    except Exception as e:

        return HttpResponse(
            JSONRenderer().render({
                'status': False,
                'message': str(e)
            }),
            content_type='application/json',
            status=400
        )


@csrf_exempt
def product_search_suggestion(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        search = python_data.get('search')
        # print(search, 'search')
        search_query = Q()
        search_query &= Q(product_name__icontains = search, status='Active') | Q(product_name_french__icontains = search, status='Active') | Q(refpro__icontains = search, status='Active')
        count_product_list = models.ProductDetail.objects.filter(search_query).count()
        # print(count_product_list, 'count_product_list')
        product_list = models.ProductDetail.objects.filter(search_query).distinct()[:15]
        product_serializer = ProductNameSerializer(product_list, many=True).data

        res={
            'data':product_serializer
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)


@csrf_exempt
def product_list_admin(request):
    product= models.ProductDetail.objects.all()
    serializer = ProductDetailSerializer(product, many=True).data
    
    res = {
        'data': serializer
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)


@csrf_exempt
def all_product_list(request):
    python_data = JSONParser().parse(io.BytesIO(request.body))

    print("python_data--all_product_list--->", python_data)

    today = timezone.now()

    page_number = int(python_data.get('page_number', 1))
    row_size = int(python_data.get('row_data', 10))

    last_row = row_size * page_number
    first_row = last_row - row_size

    category_name = python_data.get('category_name', None)
    search_key = python_data.get('search_key', None)
    subcategory_name = python_data.get('subcategory_name', None)
    status = python_data.get('status', None)
    product_added_from = python_data.get('product_added_from', None)

    filter_condition = Q()

    # Category Filter
    if category_name not in [None, '', 'null']:
        filter_condition &= Q(category__category=category_name)

    # Subcategory Filter
    if subcategory_name not in [None, '', 'null']:
        filter_condition &= Q(subcategory__subcategory=subcategory_name)

    # Search Filter
    # if search_key not in [None, '', 'null']:
    #     filter_condition &= (
    #         Q(product_name__icontains=search_key) |
    #         Q(refpro__icontains=search_key) |
    #         Q(product_code__icontains=search_key) |
    #         Q(product_name_french__icontains=search_key)
    #     )
    
    # Search Filter
    if search_key not in [None, '', 'null']:

        search_words = search_key.strip().split()

        search_query = Q()

        for word in search_words:

            search_query &= (
                Q(product_name__unaccent__icontains=word) |
                Q(refpro__unaccent__icontains=word) |
                Q(product_code__unaccent__icontains=word) |
                Q(product_name_french__unaccent__icontains=word)
            )

        filter_condition &= search_query


    # Status Filter
    if status not in [None, '', 'null']:
        filter_condition &= Q(status=status)

    # Product Added From Filter
    if product_added_from not in [None, '', 'null']:
        filter_condition &= Q(product_added_from=product_added_from)

    # Counts
    total_bulk_records = models.ProductDetail.objects.filter(
        filter_condition,
        status='bulk'
    ).count()

    total_records = models.ProductDetail.objects.filter(
        filter_condition,
        product_verification='Approved'
    ).exclude(status='bulk').count()

    total_pending_records = models.ProductDetail.objects.filter(
        product_verification='Pending'
    ).count()

    pending_inquiry_count = models.ProductInquiry.objects.filter(
        status="Pending"
    ).count()

    # Product List
    all_product = models.ProductDetail.objects.filter(
        filter_condition,
        product_verification='Approved'
    ).exclude(
        status='bulk'
    ).order_by('-id')[first_row:last_row]

    list_data = []

    for product in all_product:

        # Aggregate Min/Max Price
        price_data = models.ProductModelVariant.objects.filter(
            product=product
        ).aggregate(
            max_price=Max('price'),
            min_price=Min('price')
        )

        serializer = ProductDetailSerializer(product).data

        serializer.update({
            'max_price': price_data['max_price'],
            'min_price': price_data['min_price']
        })

        list_data.append(serializer)

    # Response
    res = {
        'data': list_data,
        'total_records': total_records,
        'current_page': page_number,
        'total_pages': int(np.ceil(total_records / row_size)),
        'total_pending_records': total_pending_records,
        'pending_inquiry_count': pending_inquiry_count,
        'total_bulk_records': total_bulk_records
    }

    json_data = JSONRenderer().render(res)

    return HttpResponse(
        json_data,
        content_type='application/json',
        status=200
    )


@csrf_exempt
def product_delete(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))

        id = python_data.get('id')

        get_product = models.ProductDetail.objects.get(id=id)
        get_product.delete()

        res={
            'message':'Product Deleted Successfully'
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)


@csrf_exempt
def product_status_update(request):
    if request.method =="POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))

        product = python_data.get('id')
        status = python_data.get('status')

        check_product = models.ProductDetail.objects.filter(id= product).count()
        if check_product == 1:
           
            update_product = models.ProductDetail.objects.get(id= product)
            update_product.status = status
            update_product.save()

            if status == 'Inactive':
                cart_delete = models.CartDetail.objects.filter(product= product)
                cart_delete.delete()

            

            res={
                'message':'Product Status update Succesfully.' 
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)

        else:
            res={
                'message':'Something went wrong.' 
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 406)


@csrf_exempt
def product_filter_data(request):
    if request.method =="POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))

        category = python_data.get('category')
        min_price = float(python_data.get('min_price'))
        max_price = float(python_data.get('max_price'))


        product_list = models.ProductDetail.objects.filter(category= category, price__gte = min_price, price__lte = max_price, status= 'Active')
        product_serializer = ProductDataSerializer(product_list, many=True).data
        

        res={
            'data':product_serializer 
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)



@csrf_exempt
def product_country_update(request):
    if request.method == "POST":
        all_product = models.ProductDetail.objects.all()
        available_country = [1, 8]
        
        for product in all_product:
            print(product.id, 'product')
            product = models.ProductDetail.objects.get(id = product.id)


            product.available_country.set(available_country)

        res = {
            'data':"serialser"
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)


@csrf_exempt
def all_product_country(request):
    # print("PRint check")
    if request.method == "POST":
        try:
            python_data = JSONParser().parse(io.BytesIO(request.body))
            countryCode = python_data.get('countryCode',None)
            # print("python_data--->",python_data)
            # print("countryCode--->",countryCode)
            if countryCode in [None,'','null']:
                country_list = models.ProductCountry.objects.annotate(
                    all_first=Case(
                        When(name='All', then=Value(0)),
                        default=Value(1),
                        output_field=IntegerField(),
                    )
                ).order_by('all_first', 'name')
            else:
                country_list = models.ProductCountry.objects.filter(Q(name = 'All')|Q(countryCode = countryCode)).order_by('id')

        except:

            country_list = models.ProductCountry.objects.annotate(
                all_first=Case(   
                    When(name='All', then=Value(0)),
                    default=Value(1),
                    output_field=IntegerField(),
                )
            ).order_by('all_first', 'name')
        
        serialser = ProductCountrySerializer(country_list, many=True).data

        # print("country_list------>",country_list)
        # print("serialser------>",serialser)
                
        res = {
            'data':serialser
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)


@csrf_exempt
def vendor_apply_product_price(request):
    python_data={}
    for i,j in request.FILES.items():
        python_data.update({i:j})
    
    for i,j in request.POST.items():
        python_data.update({i:j})
        
    # print(python_data, 'python_data')

    if python_data.get('variant_id'):
        variant_id = python_data.get('variant_id')
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')
        get_variant = models.ProductModelVariant.objects.get(id = variant_id)
        model_id = get_variant.model.id
        product_id = get_variant.product.id

        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Inactive',
        ).save()

        res = {
            'message': "Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)
    
    if python_data.get('model_id'):
        model_id = python_data.get('model_id')
        name = python_data.get('name')
        image = python_data.get('image')
        price = python_data.get('price')
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')
        name_french = python_data.get('variant_name_french')

        create_variant = models.ProductModelVariant.objects.create(
            model_id = model_id,
            name = name,
            name_french = name_french,
            image = image,
            price = price,
        )
        create_variant.save()
        variant_id = create_variant.id
        
        get_model = models.ProductModel.objects.get(id = model_id)
        product_id = get_model.product.id

        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Inactive',
        ).save()

        res = {
            'message': "Variant and Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)
    
    if python_data.get('product_id'):

        product_id = python_data.get('product_id')
        name = python_data.get('name')
        image = python_data.get('image')
        
        variant_name = python_data.get('variant_name')
        variant_image = python_data.get('variant_image')
        variant_price = python_data.get('variant_price')
        model_name_french = python_data.get('model_name_french')
        name_french = python_data.get('variant_name_french')
        
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')

        create_model = models.ProductModel.objects.create(
            product_id = product_id,
            model_name = name,
            model_name_french = model_name_french,
            model_image = image,
            status = 'Inactive'
        )
        create_model.save()

        model_id = create_model.id

        create_variant = models.ProductModelVariant.objects.create(
            model_id = model_id,
            name = variant_name,
            name_french = name_french,
            image = variant_image,
            price = variant_price,
        )
        create_variant.save()
        variant_id = create_variant.id
        
        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Inactive',
        ).save()

        res = {
            'message': "Model, Variant and Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)


    res = {
        'message': "Something went wrong."
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=406)


@csrf_exempt
def vendor_product_price_create(request):
    python_data={}
    for i,j in request.FILES.items():
        python_data.update({i:j})
    
    for i,j in request.POST.items():
        python_data.update({i:j})
        
    # print(python_data, 'python_data')

    if python_data.get('variant_id'):
        variant_id = python_data.get('variant_id')
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')
        get_variant = models.ProductModelVariant.objects.get(id = variant_id)
        model_id = get_variant.model.id
        product_id = get_variant.product.id

        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Active',
        ).save()

        res = {
            'message': "Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)
    
    if python_data.get('model_id'):
        model_id = python_data.get('model_id')
        name = python_data.get('name')
        image = python_data.get('image')
        price = python_data.get('price')
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')
        name_french = python_data.get('variant_name_french')

        create_variant = models.ProductModelVariant.objects.create(
            model_id = model_id,
            name = name,
            name_french = name_french,
            image = image,
            price = price,
        )
        create_variant.save()
        variant_id = create_variant.id
        
        get_model = models.ProductModel.objects.get(id = model_id)
        product_id = get_model.product.id

        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Active',
        ).save()

        res = {
            'message': "Variant and Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)
    
    if python_data.get('product_id'):

        product_id = python_data.get('product_id')
        name = python_data.get('name')
        image = python_data.get('image')
        
        variant_name = python_data.get('variant_name')
        variant_image = python_data.get('variant_image')
        variant_price = python_data.get('variant_price')
        model_name_french = python_data.get('model_name_french')
        name_french = python_data.get('variant_name_french')
        
        vendor = python_data.get('vendor_id')
        vendor_price = python_data.get('vendor_price')

        create_model = models.ProductModel.objects.create(
            product_id = product_id,
            model_name = name,
            model_name_french = model_name_french,
            model_image = image,
            status = 'Active'
        )
        create_model.save()

        model_id = create_model.id

        create_variant = models.ProductModelVariant.objects.create(
            model_id = model_id,
            name = variant_name,
            name_french = name_french,
            image = variant_image,
            price = variant_price,
        )
        create_variant.save()
        variant_id = create_variant.id
        
        create = models.VendorProductPrice.objects.create(
            product_id = product_id,
            model_id = model_id,
            vendor_id = vendor,
            variant_id = variant_id,
            price = vendor_price,
            status = 'Active',
        ).save()

        res = {
            'message': "Model, Variant and Vendor price added succesfully."
        }
        json_data = JSONRenderer().render(res)
        return HttpResponse(json_data, content_type= 'application/json', status=200)


    res = {
        'message': "Something went wrong."
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=406)

@csrf_exempt
def product_tag_create(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        tag = python_data.get('tag')
        category = python_data.get('category')
        if models.ProductTag.objects.filter(category = category, tag = tag).count() == 0:
            create = models.ProductTag.objects.create(
                category_id = category,
                tag = tag
            ).save()
            res = {
                'message':'Tag Added succesfully.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)
        else:
            res = {
                'message':'Something Went Wrong.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 406)


@csrf_exempt
def product_tag_update(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        id = python_data.get('id')
        category = python_data.get('category')
        tag = python_data.get('tag')
        
        if models.ProductTag.objects.filter(category=category,tag =tag, id = id).count() == 1:
            res = {
                'message':'Tags update succesfully.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)
        elif models.ProductTag.objects.filter(category=category,tag =tag).count() != 0:
            res = {
                'message':'Tags already added.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 406)
        else:
            tag_count = models.ProductTag.objects.filter(category=category,tag =tag).count()
            if tag_count == 0:
                tag_update = models.ProductTag.objects.get(id = id)
                tag_update.category_id = category 
                tag_update.tag = tag
                tag_update.save()
                res = {
                    'message':'Tags update succesfully.'
                }
                return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)
            else:
                res = {
                    'message':'Something Went Wrong.'
                }
                return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 406)


@csrf_exempt
def product_tag_delete(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        id = python_data.get('id')

        print("models.ProductTag.objects.filter(id = id).count()------>",models.ProductTag.objects.filter(id = id).count())

        if models.ProductTag.objects.filter(id = id).count() == 1:
            tag = models.ProductTag.objects.get(id = id)
            tag.delete()
            res = {
                'message':'Tag Delete succesfully.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)
        else:
            res = {
                'message':'Something Went Wrong.'
            }
            return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 406)


@csrf_exempt
def product_tag_list(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        category = python_data.get('category')

        all_product = models.ProductTag.objects.filter(category = category)
        serialiser = ProductTagSerializer(all_product, many=True).data
        
        res = {
            'data':serialiser
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)


def serialize_product(product, similarity=None):

    data = {
        "id": product.id,
        "product_name": product.product_name,
        "product_name_french": product.product_name_french,
        "category": (
            product.category.category
            if product.category else ""
        ),
        "subcategory": (
            product.subcategory.subcategory
            if product.subcategory else ""
        ),
        "color": (product.color or "").lower(),
        "product_image_1": (
            product.product_image_1.url
            if product.product_image_1
            else ""
        )
    }

    if similarity is not None:
        data["similarity"] = round(
            similarity,
            3
        )

    return data

def build_fallback_products(products):

    return [
        serialize_product(p)
        for p in products[:10]
    ]


def get_country(data):

    currency_code = data.get("currencyCode")
    calling_code = data.get("countryCallingCode")
    country = data.get("countryName")
    user_id = data.get("user_id")
    ip_address = data.get("ip_address")

    DEFAULT_COUNTRY = "Egypt"

    try:

        if currency_code not in [None, "", "null"]:

            obj = (
                models.CountryWithCurrency.objects
                .filter(currency_code=currency_code)
                .first()
            )

            if obj:
                return obj.country_name

        if calling_code not in [None, "", "null"]:

            obj = (
                models.CountryWithCurrency.objects
                .filter(
                    country_calling_code=calling_code
                )
                .first()
            )

            if obj:
                return obj.country_name

        if user_id:

            user = (
                models.CustomerDetail.objects
                .filter(id=user_id)
                .first()
            )

            if user:

                obj = (
                    models.CountryWithCurrency.objects
                    .filter(
                        country_calling_code=user.countryCode
                    )
                    .first()
                )

                if obj:
                    return obj.country_name

        if country not in [None, "", "null"]:
            return country

        if ip_address:

            response = requests.get(
                f"https://ipapi.co/{ip_address}/json/",
                timeout=3
            ).json()

            return response.get(
                "country_name",
                DEFAULT_COUNTRY
            )

    except Exception as e:
        print("Country Error =====>", e)

    return DEFAULT_COUNTRY




def fetch_products_by_ids(product_ids):

    queryset = (
        models.ProductDetail.objects
        .filter(
            id__in=product_ids,
            status="Active"
        )
        .select_related(
            "category",
            "subcategory"
        )
        .only(
            "id",
            "product_name",
            "product_name_french",
            "color",
            "product_image_1",
            "category__category",
            "subcategory__subcategory"
        )
    )

    return {
        p.id: p
        for p in queryset
    }


def build_similarity_results(
    results,
    products_map
):

    final = []

    for r in results:

        product = products_map.get(r["id"])

        if not product:
            continue

        similarity = round(
            r["similarity"],
            3
        )

        # Skip almost identical match
        if similarity > 0.995:
            continue

        final.append({
            "id": product.id,
            "product_name": product.product_name,
            "product_name_french": product.product_name_french,
            "category": (
                product.category.category
                if product.category else ""
            ),
            "subcategory": (
                product.subcategory.subcategory
                if product.subcategory else ""
            ),
            "color": (
                product.color or ""
            ).lower(),
            "similarity": similarity,
            "product_image_1": (
                product.product_image_1.url
                if product.product_image_1
                else ""
            )
        })

    return final


def apply_ranking_strategy(products):

    if not products:
        return []

    # Base similarity sorting
    products.sort(
        key=lambda x: x["similarity"],
        reverse=True
    )

    # Detect dominant subcategory
    subcategories = [
        p["subcategory"]
        for p in products
        if p["subcategory"]
    ]

    dominant_subcategory = None

    if subcategories:

        dominant_subcategory = (
            Counter(subcategories)
            .most_common(1)[0][0]
        )

    ranked_products = []

    for p in products:

        score = p["similarity"]

        # Small boost for dominant subcategory
        if (
            dominant_subcategory
            and p["subcategory"]
            == dominant_subcategory
        ):
            score += 0.03

        # Small boost if color exists
        if p["color"]:
            score += 0.01

        p["final_score"] = round(score, 4)

        ranked_products.append(p)

    ranked_products.sort(
        key=lambda x: x["final_score"],
        reverse=True
    )

    return ranked_products[:30]


def deduplicate_products(products):

    seen = set()

    return [
        p for p in products
        if not (
            p["id"] in seen
            or seen.add(p["id"])
        )
    ]

def rank_products(results):

    product_ids = [
        r["id"]
        for r in results
    ]

    products_map = fetch_products_by_ids(
        product_ids
    )

    similarity_results = (
        build_similarity_results(
            results,
            products_map
        )
    )

    ranked = apply_ranking_strategy(
        similarity_results
    )

    return deduplicate_products(
        ranked
    )




def search_similar_products(query_vector, index, ids, top_k=10):
    faiss.normalize_L2(query_vector)
    D, I = index.search(query_vector, top_k)
    results = []
    for sim, idx in zip(D[0], I[0]):
        if idx < len(ids) and sim >= 0.75:
            results.append({"id": ids[idx], "similarity": float(sim)})
    return results



@csrf_exempt
def recommend_products(request):

    if request.method != "POST":

        return JsonResponse(
            {"error": "Invalid request method"},
            status=405
        )

    try:

        python_data={}
        for i,j in request.FILES.items():
            python_data.update({i:j})

        for i,j in request.POST.items():
            python_data.update({i:j})

        country = get_country(
            python_data
        )

        image_file = python_data.get(
            "image"
        )

        if not image_file:

            return JsonResponse(
                {"error": "Image is required"},
                status=400
            )

     

        query_vector = (
            extract_features_from_file(
                image_file
            )
        )

        if (
            query_vector is None
            or query_vector.size == 0
        ):

            return JsonResponse(
                {
                    "error":
                    "Failed to extract image features"
                },
                status=400
            )

        query_vector = np.asarray(
            query_vector,
            dtype="float32"
        ).reshape(1, -1)

        # IMPORTANT NORMALIZATION
        query_vector = (
            query_vector
            / np.linalg.norm(
                query_vector,
                axis=1,
                keepdims=True
            )
        )

    

        index, ids = get_faiss_index()

        if index is None or ids is None:

            return JsonResponse(
                {
                    "error":
                    "FAISS index unavailable"
                },
                status=500
            )

   

        results = search_similar_products(
            query_vector=query_vector,
            index=index,
            ids=ids,
            top_k=50
        )

        print(
            "FAISS Results =====>",
            results
        )

        if not results:

            fallback_products = list(
                models.ProductDetail.objects
                .filter(status="Active")
                .select_related(
                    "category",
                    "subcategory"
                )
                .only(
                    "id",
                    "product_name",
                    "product_name_french",
                    "color",
                    "product_image_1",
                    "category__category",
                    "subcategory__subcategory"
                )[:10]
            )

            return JsonResponse({
                "country": country,
                "predicted_subcategory": "",
                "similar_products":
                    build_fallback_products(
                        fallback_products
                    )
            })

  
        final_results = rank_products(
            results
        )


        if not final_results:

            fallback_products = list(
                models.ProductDetail.objects
                .filter(status="Active")
                .select_related(
                    "category",
                    "subcategory"
                )
                .only(
                    "id",
                    "product_name",
                    "product_name_french",
                    "color",
                    "product_image_1",
                    "category__category",
                    "subcategory__subcategory"
                )[:10]
            )

            return JsonResponse({
                "country": country,
                "predicted_subcategory": "",
                "similar_products":
                    build_fallback_products(
                        fallback_products
                    )
            })

        dominant_subcategory = ""

        subcategories = [
            p["subcategory"]
            for p in final_results
            if p["subcategory"]
        ]

        if subcategories:

            dominant_subcategory = (
                Counter(subcategories)
                .most_common(1)[0][0]
            )

    
        return JsonResponse({

            "country": country,

            "predicted_subcategory":
                dominant_subcategory,

            "similar_products":
                final_results

        })

    except Exception as e:

        print(
            "recommend_products ERROR =====>",
            str(e)
        )

        return JsonResponse(
            {"error": str(e)},
            status=500
        )


@csrf_exempt
def recently_view_product_list(request):
    if request.method =="POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))

        customer = python_data.get('customer')

        currencyCode = python_data.get('currencyCode',None)
        countryCallingCode = python_data.get('countryCallingCode',None)
        country = python_data.get('countryName','Egypt')
        ip_address = python_data.get('ip_address')
        user_id = python_data.get('customer')

        try:
            if currencyCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(currency_code = currencyCode).first().country_name
            elif countryCallingCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(country_calling_code = countryCallingCode).first().country_name
                
            elif user_id:
                user = models.CustomerDetail.objects.get(id = user_id)
                countryCode = user.countryCode

                if models.CountryWithCurrency.objects.filter(country_calling_code = countryCode).exists():
                    country = models.CountryWithCurrency.objects.get(country_calling_code = countryCode).country_name
                else:
                    if country not in [None,'','null']:
                        country = country
                    else:
                        try:
                            response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                            response = response.json()
                            country = response.get("country_name")
                            print("response====>",response)

                            if country in [None,'','null']:
                                country = "Egypt"

                        except Exception as e:
                            print("Errroooorrr----->",e)
                            country = "Egypt"
                        

            else:
                if country in [None,'','null']:
                    try:
                        response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                        response = response.json()
                        country = response.get("country_name")
                        print("response====>",response)

                        if country in [None,'','null']:
                            country = "Egypt"

                    except Exception as e:
                        print("Errroooorrr----->",e)
                        country = "Egypt"
                else:
                    country = country
        except Exception as e:
            print("Error-=-=-=--->",e)
            country = 'Egypt'

        product = models.RecentlyViewProduct.objects.filter(customer = customer).order_by('-id')
        serializer = RecentlyViewProductSerializer(product, many=True, context = {'country':country}).data

        res={
            'data':serializer
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)
    


@csrf_exempt
def product_remove_add_to_cart(request):
    if request.method == "POST":
        removed_carted_products = models.WishlistDetail.objects.filter(product__isnull=True)
        res={
            "removed_carted_products":removed_carted_products
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)




# @csrf_exempt
# def bulk_upload(request):
#     if (request.method == "POST"):
#         python_data={}
#         for i,j in request.FILES.items():
#             python_data.update({i:j})
        
#         for i,j in request.POST.items():
#             python_data.update({i:j})
        
#         excel_file = python_data.get('excel_file')
#         flag = python_data.get('flag')

        
#         letters = ['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U','V','W','X','Y','Z',\
#                 'AA','AB','AC','AD','AE','AF','AG','AH','AI','AJ','AK','AL','AM','AN','AO','AP','AQ','AR','AS','AT','AU','AV','AW','AX','AY','AZ', \
#                 'BA','BB','BC','BD','BE','BF','BG','BH','BI','BJ','BK','BL','BM','BN','BO','BP','BQ','BR','BS','BT','BU','BV','BW','BX','BY','BZ']
            
#         # all_subcategory_list = models.SubCategoryDetail.objects.all()
#         # for get_subcategory in all_subcategory_list:
#         #     get_subcategory.subcategory = get_subcategory.subcategory.strip()
#         #     get_subcategory.subcategory_french = get_subcategory.subcategory_french.strip()
#         #     get_subcategory.save()

#         try:
#             fs = FileSystemStorage()
#             fs.delete('product_import.xlsx')
#         except:
#             pass
            
#         fs = FileSystemStorage()
#         filename = fs.save('product_import.xlsx', excel_file)
#         file_url = fs.url(filename)
#         file_path = fs.path(filename)  
#         workbook = load_workbook(file_path)
#         sheet = workbook.active
#         invalid_count = 0

#         new_filename = 'product_import_data.xlsx'
#         new_file_url = fs.url(new_filename)
#         new_file_path = fs.path(new_filename)  

#         wb = load_workbook(new_file_path)
#         ws = wb.active
#         ws.delete_rows(2, ws.max_row) 
#         wb.save(new_file_path)
        
#         new_workbook = load_workbook(new_file_path)
#         new_sheet = new_workbook.active

#         for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
#             if row_index == 1:
#                 continue
            
#             if row[0] == None:
#                 break
            
#             print(row[0], 'row')
                        
#             category = 	(row[0]).strip()
#             subcategory = (row[1]).strip()
#             unit_of_measure = (row[2]).lower()
#             shipping_via = row[3]
#             country_of_origin = (row[4]).strip()
#             available_country = [v.strip() for v in row[5].split(",")]
#             tags = [t.strip() for t in row[6].split(",")]
#             product_name = 	row[7]
#             description = 	row[8]
#             price = row[9]            
#             final_price = row[10]
#             min_order_quantity = row[11]
#             length = row[12]
#             width = row[13]
#             height = row[14]
#             weight = row[15]
#             vendor = row[16]
#             reuser =row[17]
#             refpro =row[18]
#             product_type =row[19]
#             product_packaging =row[20]
#             product_packaging_value =row[21]

#             print(tags, 'tags')
            

#             get_new_cell = str(letters[0]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[0]  
#             new_workbook.save(new_file_path)
    
#             get_new_cell = str(letters[1]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[1]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[2]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[2]
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[3]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[3]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[4]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[4]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[5]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[5]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[6]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[6]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[7]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[7]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[8]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[8]  
#             new_workbook.save(new_file_path)

#             get_new_cell = str(letters[9]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[9]  
#             new_workbook.save(new_file_path)


#             get_new_cell = str(letters[10]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[10]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[11]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[11]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[12]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[12]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[13]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[13]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[14]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[14]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[15]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[15]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[16]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[16]  
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[17]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[17]  
#             new_workbook.save(new_file_path)


#             get_new_cell = str(letters[18]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[18]  
#             new_workbook.save(new_file_path)

            
#             get_new_cell = str(letters[19]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[19]  
#             new_workbook.save(new_file_path)

#             get_new_cell = str(letters[20]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[20]  
#             new_workbook.save(new_file_path)

#             get_new_cell = str(letters[21]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = row[21]  
#             new_workbook.save(new_file_path)


#             get_new_cell = str(letters[23]) + str(1)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = 'Status' 
#             new_workbook.save(new_file_path)
            
#             get_new_cell = str(letters[24]) + str(1)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = 'Reasons' 
#             new_workbook.save(new_file_path)



#             category_reason = None
#             print("category---->",category)
#             print("subcategory---->",subcategory)
#             check_category = models.CategoryDetail.objects.filter(Q(category__iexact = category)| Q(category_french__iexact = category)).count()
#             # print(check_category, 'check_category')
#             if check_category == 0:
#                 get_cell = str(letters[0]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 category_reason = 'Enter Valid Category'
#                 invalid_count += 1
#                 category_id = None
#             else:
#                 category_id = models.CategoryDetail.objects.get(Q(category__iexact = category)| Q(category_french__iexact = category)).id
                
#             print(category_id, 'category_id')
            
#             sub_category_reason = None
#             check_subcategory = models.SubCategoryDetail.objects.filter(Q(subcategory__iexact =subcategory) | Q(subcategory_french__iexact = subcategory)).count()
#             # print(check_subcategory, 'check_subcategory')
#             if check_subcategory == 1:
#                 check_category_subcategory = models.SubCategoryDetail.objects.annotate(
#                     category_clean=Trim('category__category'),
#                     category_fr_clean=Trim('category__category_french'),
#                     subcategory_clean=Trim('subcategory'),
#                     subcategory_fr_clean=Trim('subcategory_french'),
#                 ).filter(
#                     (Q(category_clean__iexact=category) | Q(category_fr_clean__iexact=category)) &
#                     (Q(subcategory_clean__iexact=subcategory) | Q(subcategory_fr_clean__iexact=subcategory))
#                 ).count()                
                
#                 if check_category_subcategory == 0:
#                     get_cell = str(letters[1]) + str(row_index)
#                     cell_bg = new_sheet[get_cell]
#                     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                     new_workbook.save(new_file_path)
#                     sub_category_reason = 'This Subcategory not associate with this category'
#                     invalid_count += 1
#             else:
#                 get_cell = str(letters[1]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 sub_category_reason = 'Enter Valid Subcategory'
#                 invalid_count += 1
            
                        
#             unit_of_measure_reason = None
#             if unit_of_measure not in ['kg', 'pcs', 'liters', 'meters', 'tons', 'carton', 'douzaine', "1/2douzaine"]:
#                 get_cell = str(letters[2]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 unit_of_measure_reason = 'Please select one value from the following options: kg, pcs, liters, meters, tons, carton, douzaine, or 1/2 douzaine.'
#                 invalid_count += 1
            


#             shipping_via_reason = None
#             if shipping_via.lower() not in ['By Air', 'by air', 'By Ship', 'by ship', 'Both', 'both']:
#                 get_cell = str(letters[3]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 shipping_via_reason = 'Please select one value from the following options: By Air, By Ship, or Both.'
#                 invalid_count += 1


            

#             country_of_origin_reason = None
#             check_country_of_origin = models.Country.objects.filter(country_name__iexact = country_of_origin).count()
#             # print(check_category, 'check_category')
#             if check_country_of_origin == 0:
#                 get_cell = str(letters[4]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 country_of_origin_reason = 'Enter Valid country of origin'
#                 invalid_count += 1


#             available_country_reason = None
#             list_not_list_country = []

#             available_country = [c.strip() for c in available_country if c.strip()]

#             has_all = any(c.lower() == 'all' for c in available_country)

#             if has_all:
#                 get_new_cell = str(letters[5]) + str(row_index)
#                 cell_new_bg = new_sheet[get_new_cell]
#                 cell_new_bg.value = 'All'
#             else:
#                 for country in available_country:
#                     check_country = models.ProductCountry.objects.filter(Q(name__iexact=country) | Q(name_french__iexact=country)).exists()
#                     if not check_country:
#                         print(country, 'country')
#                         list_not_list_country.append(country)

#                 if list_not_list_country:
#                     get_cell = str(letters[5]) + str(row_index)
#                     cell_bg = new_sheet[get_cell]
#                     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                     available_country_reason = f'{list_not_list_country} country not mention in system'
#                     invalid_count += 1

#             for tag in tags:
#                 check_tag = models.ProductTag.objects.filter(category_id =category_id, tag = tag ).exists()
#                 if check_tag == 0:
#                     check_tag = models.ProductTag.objects.create(
#                         category_id =category_id, 
#                         tag = tag 
#                     ).save()

#             price_reason = None
#             price_str = str(price).strip()

#             get_cell = str(letters[9]) + str(row_index)
#             cell_bg = new_sheet[get_cell]

#             if price_str.startswith("="):
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 price_reason = 'Formula not allowed. Enter numeric value'
#                 invalid_count += 1

#             else:
#                 try:
#                     # Remove commas and convert
#                     price_value = float(price_str.replace(",", ""))

#                     if price_value < 0:
#                         cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                         price_reason = 'Price cannot be negative'
#                         invalid_count += 1

#                     else:
#                         # Round if needed
#                         price_value = round(price_value)

#                         new_sheet[get_cell].value = price_value
#                         new_workbook.save(new_file_path)

#                 except ValueError:
#                     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                     price_reason = 'Enter numeric value only'
#                     invalid_count += 1


#             final_price_reason = None
#             final_price_str = str(final_price).strip()

#             get_cell = str(letters[10]) + str(row_index)
#             cell_bg = new_sheet[get_cell]

#             if final_price_str.startswith("="):
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 final_price_reason = 'Formula not allowed. Enter numeric value'
#                 invalid_count += 1

#             else:
#                 try:
#                     # Convert safely
#                     final_price_value = float(final_price_str.replace(",", ""))

#                     if final_price_value < 0:
#                         cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                         final_price_reason = 'Price cannot be negative'
#                         invalid_count += 1

#                     else:
#                         final_price_value = round(final_price_value)

#                         new_sheet[get_cell].value = final_price_value
#                         new_workbook.save(new_file_path)

#                 except ValueError:
#                     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                     final_price_reason = 'Enter numeric value only'
#                     invalid_count += 1

#             # min_order_quantity_reason = None
#             # if int(min_order_quantity) < 1:
#             #     get_cell = str(letters[10]) + str(row_index)
#             #     cell_bg = new_sheet[get_cell]
#             #     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#             #     min_order_quantity_reason = 'Min order qunatity not less than 0'
#             #     invalid_count += 1

#             min_order_quantity_reason = None
#             min_order_quantity_str = str(min_order_quantity).strip()

#             get_cell = str(letters[11]) + str(row_index)
#             cell_bg = new_sheet[get_cell]

#             if min_order_quantity_str.startswith("="):
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 min_order_quantity_reason = 'Formula not allowed in Min order quantity. Enter numeric value'
#                 invalid_count += 1

#             else:
#                 try:
#                     # Convert safely (also handles commas)
#                     min_qty_value = int(float(min_order_quantity_str.replace(",", "")))

#                     if min_qty_value < 1:
#                         cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                         min_order_quantity_reason = 'Min order quantity must be at least 1'
#                         invalid_count += 1

#                     else:
#                         new_sheet[get_cell].value = min_qty_value
#                         new_workbook.save(new_file_path)

#                 except ValueError:
#                     cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                     min_order_quantity_reason = 'Enter numeric value only'
#                     invalid_count += 1

             
#             vendor_reason = None
#             check_vendor = models.VendorDetail.objects.filter(email__iexact = vendor).count()
#             # print(check_category, 'check_category')
#             if check_vendor == 0:
#                 get_cell = str(letters[16]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 vendor_reason = 'Enter valid seller email'
#                 invalid_count += 1

#             length_reason = None
#             try:
#                 length_value = float(length)

#                 if length_value < 0:
#                     length_reason = 'Value cannot be negative'
#                 elif length_value == 0:
#                     length_reason = 'Value must be greater than 0'

#             except ValueError:
#                 length_reason = 'Enter numeric values only'

#             if length_reason:
#                 get_cell = str(letters[12]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1

            
#             width_reason = None
#             try:
#                 width_value = float(width)

#                 if width_value < 0:
#                     width_reason = 'Value cannot be negative'
#                 elif width_value == 0:
#                     width_reason = 'Value must be greater than 0'

#             except ValueError:
#                 width_reason = 'Enter numeric values only'

#             if width_reason:
#                 get_cell = str(letters[13]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1
            

#             height_reason = None
#             try:
#                 height_value = float(height)

#                 if height_value < 0:
#                     height_reason = 'Value cannot be negative'
#                 elif height_value == 0:
#                     height_reason = 'Value must be greater than 0'

#             except ValueError:
#                 height_reason = 'Enter numeric values only'

#             if height_reason:
#                 get_cell = str(letters[14]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1

#             weight_reason = None
#             if shipping_via.lower() == "by ship":
#                 try:
#                     weight_value = float(weight)

#                     if weight_value < 0:
#                         weight_reason = 'Value cannot be negative'
#                     elif weight_value == 0:
#                         weight_reason = 'Value must be greater than 0'

#                 except ValueError:
#                     weight_reason = 'Enter numeric values only'

#             if weight_reason:
#                 get_cell = str(letters[15]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1

#             reuser_reason = None
#             if reuser is None or str(reuser).strip() == "":
#                 reuser_reason = 'This field cannot be empty'

#             if reuser_reason:
#                 get_cell = str(letters[17]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1


#             # Refpro validation
#             refpro_reason = None
#             if refpro is None or str(refpro).strip() == "":
#                 refpro_reason = 'This field cannot be empty'

#             if refpro_reason:
#                 get_cell = str(letters[18]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1

#             product_type_reason = None
#             check_product_type = models.ProductType.objects.filter(product_type = product_type).count()
#             if check_product_type == 0:
#                 get_cell = str(letters[19]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 category_reason = 'Enter Valid Product Type'
#                 invalid_count += 1

#             product_packaging_reason = None
#             check_product_packaging = models.ProductPackagingBy.objects.filter(product_packaging = product_packaging).count()
#             if check_product_packaging == 0:
#                 get_cell = str(letters[20]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00",fill_type = "solid")
#                 new_workbook.save(new_file_path)
#                 category_reason = 'Enter Valid Product Packaging By'
#                 invalid_count += 1

#             product_packaging_value_reason = None

#             if isinstance(product_packaging_value, float):
#                 product_packaging_value_reason = 'Enter integer values only'

#             elif not str(product_packaging_value).isdigit():
#                 product_packaging_value_reason = 'Enter integer values only'

#             else:
#                 product_packaging_value_value = int(product_packaging_value)

#                 if product_packaging_value_value < 0:
#                     product_packaging_value_reason = 'Value cannot be negative'
#                 elif product_packaging_value_value == 0:
#                     product_packaging_value_reason = 'Value must be greater than 0'


#             if product_packaging_value_reason:
#                 get_cell = str(letters[21]) + str(row_index)
#                 cell_bg = new_sheet[get_cell]
#                 cell_bg.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
#                 new_workbook.save(new_file_path)
#                 invalid_count += 1
                        
#             list_reason = []
#             if category_reason != None:
#                 list_reason.append(category_reason)

#             if sub_category_reason != None:
#                 list_reason.append(sub_category_reason)            
            
#             if unit_of_measure_reason != None:
#                 list_reason.append(unit_of_measure_reason)            
            
#             if shipping_via_reason != None:
#                 list_reason.append(shipping_via_reason)
            
#             if country_of_origin_reason != None:
#                 list_reason.append(country_of_origin_reason)
            
#             if available_country_reason != None:
#                 list_reason.append(available_country_reason)

#             if price_reason != None:
#                 list_reason.append(price_reason)
    
#             if final_price_reason != None:
#                 list_reason.append(final_price_reason)
            
#             if min_order_quantity_reason != None:
#                 list_reason.append(min_order_quantity_reason)
            
#             if vendor_reason != None:
#                 list_reason.append(vendor_reason)
            
#             if length_reason != None:
#                 list_reason.append(length_reason)
    
#             if width_reason != None:
#                 list_reason.append(width_reason)
            
#             if height_reason != None:
#                 list_reason.append(height_reason)
            
#             if weight_reason != None:
#                 list_reason.append(weight_reason)    
            
#             if product_type_reason != None:
#                 list_reason.append(product_type_reason)    

#             if product_packaging_reason != None:
#                 list_reason.append(product_packaging_reason)    
            
#             if product_packaging_value_reason != None:
#                 list_reason.append(product_packaging_value_reason)    
    
#             if reuser_reason != None:
#                 list_reason.append(reuser_reason)    
    
#             if refpro_reason != None:
#                 list_reason.append(refpro_reason)    
    


#             # final_reason = ', '.join(map(str, list_reason)) 

#             final_reason = ', '.join(f"'{item}'" for item in list_reason)


#             # print(final_reason, 'fdinal reason check ')
#             if final_reason == '':
#                 row_status = 'Valid'
#                 start_color="3d7d13" 
#                 end_color="3d7d13"
#             else:
#                 row_status = 'Invalid'
#                 start_color="e39830"
#                 end_color="e39830"


#             get_new_cell = str(letters[23]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = str(row_status)
#             cell_new_bg.fill = PatternFill(start_color=start_color, end_color=end_color,fill_type = "solid")
#             new_workbook.save(new_file_path)


#             get_new_cell = str(letters[24]) + str(row_index)
#             cell_new_bg = new_sheet[get_new_cell]
#             new_sheet[get_new_cell].value = str(final_reason)  
#             new_workbook.save(new_file_path)

#         if invalid_count == 0:
#             flag = 'Valid'
#         else:
#             flag = 'Invalid'
#         # link = SITE_MEDIA_BASE_URL + new_filename
#         link = SITE_MEDIA_BASE_URL + new_filename
#         response_data = {
#             "statuscode":200,
#             "status":"Success",
#             'message':'File Upload successulfully.',
#             "invalid_count":invalid_count,
#             "flag":flag,
#             "link":link
#         }
#         json_data = JSONRenderer().render(response_data)
#         return HttpResponse(json_data, content_type= 'application/json')



# @csrf_exempt
# def bulk_upload_submit(request):
#     if (request.method == "POST"):

#         python_data = JSONParser().parse(io.BytesIO(request.body))

#         # flag = python_data.get('flag')
#         link = python_data.get('link')

#         now = datetime.now()
#         current_time = now.strftime("%Y-%m-%d, %H:%M:%S")
        
#         letters = ['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U','V','W','X','Y','Z',\
#                 'AA','AB','AC','AD','AE','AF','AG','AH','AI','AJ','AK','AL','AM','AN','AO','AP','AQ','AR','AS','AT','AU','AV','AW','AX','AY','AZ']
            
        
        
#         fs = FileSystemStorage()
#         file_path = fs.path(link)  
#         workbook = load_workbook(file_path)
#         sheet = workbook.active
#         invalid_count = 0
#         for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
#             if row_index == 1:
#                 continue

                                    
#             category = 	(row[0]).strip()
#             subcategory = (row[1]).strip()
#             unit_of_measure = (row[2]).lower()
#             shipping_via = row[3]
#             country_of_origin = (row[4]).strip()
#             available_country = [v.strip() for v in row[5].split(",")]
#             tags = [t.strip() for t in row[6].split(",")]
#             product_name = 	row[7]
#             description = row[8]
#             price = row[9]            
#             final_price = row[10]
#             min_order_quantity = row[11]
#             length = row[12]
#             width = row[13]
#             height = row[14]
#             weight = row[15]
#             vendor = row[16]
#             reuser =row[17]
#             refpro =row[18]
#             product_type =row[19]
#             product_packaging =row[20]
#             product_packaging_value =row[21]


#             digits = '123456789' 
#             random_number = ""
#             for i in range(3) :
#                 random_number += digits[math.floor(random.random() * 9)]



#             category_name = models.CategoryDetail.objects.filter(Q(category__iexact = category)| Q(category_french__iexact = category)).values_list('category', flat=True)[0][:3].upper()
#             subcategory_name = models.SubCategoryDetail.objects.filter(Q(subcategory__iexact = subcategory)| Q(subcategory_french__iexact = subcategory)).values_list('subcategory', flat=True)[0][:3].upper()
#             product_name_data = (product_name)[:3].upper()
#             product_code =  category_name+'_'+subcategory_name+ '_'+ product_name_data + '_' + random_number

#             category_id = models.CategoryDetail.objects.get(Q(category__iexact = category)| Q(category_french__iexact = category)).id
#             subcategory_id = models.SubCategoryDetail.objects.get(Q(subcategory__iexact = subcategory)| Q(subcategory_french__iexact = subcategory)).id
#             product_type_id = models.ProductType.objects.get(product_type__iexact = product_type).id
#             product_packaging_id = models.ProductPackagingBy.objects.get(product_packaging__iexact = product_packaging).id
            
#             country_of_origin_id = models.Country.objects.get(country_name = country_of_origin).id
#             vendor_id = models.VendorDetail.objects.get(email = vendor).id

#             list_available_country = []
#             for available in available_country:
#                 available_id = models.ProductCountry.objects.get(Q(name__iexact=available) | Q(name_french__iexact=available)).id
#                 list_available_country.append(available_id)

#             list_tags = []
#             for tag in tags:
#                 tag_id = models.ProductTag.objects.get(category =category_id,  tag__exact=tag).id
#                 list_tags.append(tag_id)
            
#             shipping_via = ' '.join(word.capitalize() for word in shipping_via.split())
#             print(shipping_via, 'shipping_via')
#             product = models.ProductDetail.objects.create(
#                 product_name = product_name,
#                 product_name_french = product_name,
#                 product_code = product_code,
#                 vendor_id = vendor_id,
#                 category_id = category_id,
#                 subcategory_id = subcategory_id,
#                 description = description,
#                 description_french = description,
#                 unit_of_measure = unit_of_measure,
#                 length = length,
#                 height = height,
#                 width = width,
#                 weight = weight,
#                 country_of_origin_id = country_of_origin_id,
#                 carton_length = length,
#                 carton_width = width,
#                 carton_height = height,
#                 carton_weight = weight,
#                 product_verification = 'Approved',
#                 min_order_quantity = min_order_quantity,
#                 price = price,
#                 discount = 0,
#                 final_price = final_price,
#                 status = 'bulk',
#                 # status = 'Active',
#                 product_added_from = 'bulk',
#                 refpro = refpro,
#                 reuser = reuser,
#                 product_type_id = product_type_id,
#                 product_packaging_id = product_packaging_id,
#                 product_packaging_value = product_packaging_value,
#                 shipping_via = shipping_via,
#             )
#             product.save()
#             try: 
#                 product.available_country.set(list_available_country)
#             except:
#                 product.available_country.clear()

#             try: 
#                 product.tag.set(list_tags)
#             except:
#                 product.tag.clear()

#             product_id = product.id 
#             create_model = models.ProductModel.objects.create(
#                 product_id = product_id,
#                 model_name = product_name,
#                 model_name_french = product_name,
#                 # model_image = model_image,
#                 status = 'Active'
#             )
#             create_model.save()
#             list_price = []
            
#             create_variant = models.ProductModelVariant.objects.create(
#                 product_id = product_id,
#                 model_id = create_model.id,
#                 name = product_name,
#                 name_french = product_name,
#                 # image = variant_image,
#                 price = final_price,
#                 variant_verification = 'Approved',
#                 status = 'Active'

#             )
#             create_variant.save()
#             variant_id = create_variant.id
         
#             price = models.VendorProductPrice.objects.create(
#                 product_id = product_id,
#                 model_id = create_model.id,
#                 vendor_id = vendor_id,
#                 variant_id = variant_id,
#                 price = price,
#                 # quantity = quantity,
#                 status = 'Active'
#             ).save()

#             smallest_price = final_price
#             product_update = models.ProductDetail.objects.get(id = product_id)
#             product_update.price = smallest_price
#             product_update.save()
            
#         response_data = {
#             "statuscode":200,
#             "status":"Success",
#             'message':'Data submitted successfully.',
#         }
#         json_data = JSONRenderer().render(response_data)
#         return HttpResponse(json_data, content_type= 'application/json')

############## S3 File Store start here ##############
#####          DO NOT DELETE THIS CODE           #####
######################################################

from openpyxl import Workbook

@csrf_exempt
def bulk_upload(request):
    if request.method != "POST":
        return HttpResponse(status=405)


    python_data = {}
    for i, j in request.FILES.items():
        python_data[i] = j
    for i, j in request.POST.items():
        python_data[i] = j

    excel_file = python_data.get('excel_file')
    flag = python_data.get('flag')

    letters = ['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U','V','W','X','Y','Z',
               'AA','AB','AC','AD','AE','AF','AG','AH','AI','AJ','AK','AL','AM','AN','AO','AP','AQ','AR','AS','AT','AU','AV','AW','AX','AY','AZ',
               'BA','BB','BC','BD','BE','BF','BG','BH','BI','BJ','BK','BL','BM','BN','BO','BP','BQ','BR','BS','BT','BU','BV','BW','BX','BY','BZ']

    # -------- SAVE INPUT FILE TO S3 -------- #
    if fs.exists('product_import.xlsx'):
        fs.delete('product_import.xlsx')

    filename = fs.save('product_import.xlsx', excel_file)

    with fs.open(filename, 'rb') as f:
        workbook = load_workbook(f)

    sheet = workbook.active
    invalid_count = 0
    new_workbook = Workbook()
    new_sheet = new_workbook.active

    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):

        if row_index == 1:
            for col_index in range(len(row)):
                new_sheet[f"{letters[col_index]}1"] = row[col_index]

            # Add your extra headers
            new_sheet[f"{letters[23]}1"] = "Status"
            new_sheet[f"{letters[24]}1"] = "Reasons"
            continue

        if row[0] is None:
            break

        # -------- READ VALUES -------- #
        category = str(row[0]).strip()
        subcategory = str(row[1]).strip()
        unit_of_measure = str(row[2]).lower()
        shipping_via = row[3]
        country_of_origin = str(row[4]).strip()
        available_country = [v.strip() for v in str(row[5]).split(",")]
        tags = [t.strip() for t in str(row[6]).split(",")]

        product_name = row[7]
        description = row[8]
        price = row[9]
        final_price = row[10]
        min_order_quantity = row[11]

        length = row[12]
        width = row[13]
        height = row[14]
        weight = row[15]
        vendor = row[16]
        reuser = row[17]
        refpro = row[18]
        product_type = row[19]
        product_packaging = row[20]
        product_packaging_value = row[21]

        # -------- COPY ROW -------- #
        for col_index in range(len(row)):
            new_sheet[f"{letters[col_index]}{row_index}"] = row[col_index]

        list_reason = []

        # -------- CATEGORY -------- #
        category_id = None
        if not models.CategoryDetail.objects.filter(
            Q(category__iexact=category) | Q(category_french__iexact=category)
        ).exists():
            new_sheet[f"{letters[0]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Enter Valid Category")
            invalid_count += 1
        else:
            category_id = models.CategoryDetail.objects.filter(
                Q(category__iexact=category) | Q(category_french__iexact=category)
            ).first().id

        # -------- SUBCATEGORY -------- #
        if not models.SubCategoryDetail.objects.filter(
            Q(subcategory__iexact=subcategory) | Q(subcategory_french__iexact=subcategory)
        ).exists():
            new_sheet[f"{letters[1]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Enter Valid Subcategory")
            invalid_count += 1

        # -------- UNIT -------- #
        if unit_of_measure not in ['kg','pcs','liters','meters','tons','carton','douzaine','1/2douzaine']:
            new_sheet[f"{letters[2]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid unit")
            invalid_count += 1

        # -------- SHIPPING -------- #
        if shipping_via not in ['By Air','by air','By Ship','by ship','Both','both']:
            new_sheet[f"{letters[3]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Please select one value from the following options: By Air, By Ship, or Both.")
            invalid_count += 1

        # -------- COUNTRY -------- #
        if not models.Country.objects.filter(country_name__iexact=country_of_origin).exists():
            new_sheet[f"{letters[4]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid country")
            invalid_count += 1

        # -------- AVAILABLE COUNTRY -------- #
#         for country in available_country:
#             if country.lower() != 'all':
#                 if not models.ProductCountry.objects.filter(name__iexact=country).exists():
#                     new_sheet[f"{letters[5]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
#                     list_reason.append(f"{country} not found")
#                     invalid_count += 1

        available_country = [c.strip() for c in available_country if c.strip()]

        # If 'all' exists, keep only 'All'
        if any(country.lower() == 'all' for country in available_country):
            available_country = ['All']

        for country in available_country:
            if country.lower() != 'all':
                if not models.ProductCountry.objects.filter(
                    name__iexact=country
                ).exists():

                    new_sheet[f"{letters[5]}{row_index}"].fill = PatternFill(
                        fill_type="solid",
                        start_color="FFFF00"
                    )

                    list_reason.append(f"{country} not found")
                    invalid_count += 1




        # -------- PRICE -------- #
        try:
            price_val = float(str(price).replace(",", ""))
            if price_val < 0:
                raise Exception()
            new_sheet[f"{letters[9]}{row_index}"] = round(price_val)
        except:
            new_sheet[f"{letters[9]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid price")
            invalid_count += 1

        # -------- FINAL PRICE -------- #
        try:
            val = float(str(final_price).replace(",", ""))
            if val < 0:
                raise Exception()
            new_sheet[f"{letters[10]}{row_index}"] = round(val)
        except:
            new_sheet[f"{letters[10]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid final price")
            invalid_count += 1

        # -------- MIN QTY -------- #
        try:
            val = int(float(str(min_order_quantity)))
            if val < 1:
                raise Exception()
            new_sheet[f"{letters[11]}{row_index}"] = val
        except:
            new_sheet[f"{letters[11]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid min qty")
            invalid_count += 1

        # -------- DIMENSIONS -------- #
        for idx, val in [(12,length),(13,width),(14,height)]:
            try:
                v = float(val)
                if v <= 0:
                    raise Exception()
            except:
                new_sheet[f"{letters[idx]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
                list_reason.append("Invalid dimension")
                invalid_count += 1

        if weight is None or str(weight).strip().lower() in ['null', '', '-']:
            # weight = 0
            new_sheet[f"{letters[15]}{row_index}"] =  0
        # -------- VENDOR -------- #
        if not models.VendorDetail.objects.filter(email__iexact=vendor).exists():
            new_sheet[f"{letters[16]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid vendor")
            invalid_count += 1

        # -------- REQUIRED FIELDS -------- #
        if not reuser:
            new_sheet[f"{letters[17]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Reuser empty")
            invalid_count += 1

        if not refpro:
            new_sheet[f"{letters[18]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Refpro empty")
            invalid_count += 1

        # -------- PRODUCT TYPE -------- #
        if not models.ProductType.objects.filter(product_type=product_type).exists():
            new_sheet[f"{letters[19]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid product type")
            invalid_count += 1

        # -------- PACKAGING -------- #
        if not models.ProductPackagingBy.objects.filter(product_packaging=product_packaging).exists():
            new_sheet[f"{letters[20]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid packaging")
            invalid_count += 1

        # -------- PACKAGING VALUE -------- #
        if not str(product_packaging_value).isdigit() or int(product_packaging_value) <= 0:
            new_sheet[f"{letters[21]}{row_index}"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
            list_reason.append("Invalid packaging value")
            invalid_count += 1

        # -------- STATUS -------- #
        status = "Valid" if not list_reason else "Invalid"
        color = "3d7d13" if status == "Valid" else "e39830"

        new_sheet[f"{letters[23]}{row_index}"] = status
        new_sheet[f"{letters[23]}{row_index}"].fill = PatternFill(fill_type="solid", start_color=color)

        # new_sheet[f"{letters[24]}{row_index}"] = ", ".join(list_reason)
        new_sheet[f"{letters[24]}{row_index}"] = ', '.join(f"'{item}'" for item in list_reason)

    # -------- SAVE OUTPUT TO S3 -------- #
    output = BytesIO()
    new_workbook.save(output)
    output.seek(0)

    new_filename = "product_import_data.xlsx"

    if fs.exists(new_filename):
        fs.delete(new_filename)

    fs.save(new_filename, output)

    link = SITE_MEDIA_BASE_URL + new_filename
    print(link, 'link')

    response_data = {
        "statuscode": 200,
        "status": "Success",
        "message": "File processed successfully",
        "invalid_count": invalid_count,
        "flag": "Valid" if invalid_count == 0 else "Invalid",
        "link": link
    }

    return HttpResponse(JSONRenderer().render(response_data), content_type='application/json')


@csrf_exempt
def bulk_upload_submit(request):
    if request.method != "POST":
        return HttpResponse(status=405)

    try:
        python_data = JSONParser().parse(io.BytesIO(request.body))
        link = python_data.get('link')
        link = SITE_MEDIA_BASE_URL + link


        if not link:
            return HttpResponse(
                JSONRenderer().render({"status": "Error", "message": "File link required"}),
                content_type='application/json'
            )

        # -------- READ FILE FROM S3 -------- #
        filename = os.path.basename(link)

        with fs.open(filename, 'rb') as f:

            workbook = load_workbook(f)

        sheet = workbook.active

        # -------- LOOP THROUGH EXCEL -------- #
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):

            if row_index == 1:
                continue  # skip header

            if not row or not row[0]:
                continue

            # -------- SAFE DATA EXTRACTION -------- #
            category = str(row[0] or "").strip()
            subcategory = str(row[1] or "").strip()
            unit_of_measure = str(row[2] or "").lower()
            shipping_via = str(row[3] or "").title()
            country_of_origin = str(row[4] or "").strip()

            available_country = [v.strip() for v in str(row[5] or "").split(",") if v]
            tags = [t.strip() for t in str(row[6] or "").split(",") if t]

            product_name = str(row[7] or "").strip()
            description = str(row[8] or "").strip()

            price = float(row[9] or 0)
            final_price = float(row[10] or 0)
            min_order_quantity = int(float(row[11] or 1))

            length = float(row[12] or 0)
            width = float(row[13] or 0)
            height = float(row[14] or 0)
            weight = float(row[15] or 0)

            vendor = str(row[16] or "").strip()
            reuser = row[17]
            refpro = row[18]

            product_type = str(row[19] or "").strip()
            product_packaging = str(row[20] or "").strip()
            product_packaging_value = int(row[21] or 0)

            # -------- FETCH DATABASE OBJECTS -------- #
            category_obj = models.CategoryDetail.objects.filter(
                Q(category__iexact=category) | Q(category_french__iexact=category)
            ).first()

            subcategory_obj = models.SubCategoryDetail.objects.filter(
                Q(subcategory__iexact=subcategory) | Q(subcategory_french__iexact=subcategory)
            ).first()

            product_type_obj = models.ProductType.objects.filter(
                product_type__iexact=product_type
            ).first()

            packaging_obj = models.ProductPackagingBy.objects.filter(
                product_packaging__iexact=product_packaging
            ).first()

            country_obj = models.Country.objects.filter(
                country_name__iexact=country_of_origin
            ).first()

            vendor_obj = models.VendorDetail.objects.filter(
                email__iexact=vendor
            ).first()

            # -------- SKIP INVALID ROW -------- #
            if not all([category_obj, subcategory_obj, product_type_obj, packaging_obj, country_obj, vendor_obj]):
                continue

            # -------- GENERATE PRODUCT CODE -------- #
            random_number = str(random.randint(1000, 9999))
            product_code = (
                category_obj.category[:3].upper() + "_" +
                subcategory_obj.subcategory[:3].upper() + "_" +
                product_name[:3].upper() + "_" +
                random_number
            )

            # -------- AVAILABLE COUNTRIES -------- #
            list_available_country = []
            for available in available_country:
                obj = models.ProductCountry.objects.filter(
                    Q(name__iexact=available) | Q(name_french__iexact=available)
                ).first()
                if obj:
                    list_available_country.append(obj.id)

            # -------- TAGS -------- #
            list_tags = []
            for tag in tags:
                obj = models.ProductTag.objects.filter(
                    category=category_obj.id,
                    tag__iexact=tag
                ).first()
                if obj:
                    list_tags.append(obj.id)

            # -------- CREATE PRODUCT -------- #
            product = models.ProductDetail.objects.create(
                product_name=product_name,
                product_name_french=product_name,
                product_code=product_code,
                vendor_id=vendor_obj.id,
                category_id=category_obj.id,
                subcategory_id=subcategory_obj.id,
                description=description,
                description_french=description,
                unit_of_measure=unit_of_measure,
                length=length,
                height=height,
                width=width,
                weight=weight,
                country_of_origin_id=country_obj.id,
                carton_length=length,
                carton_width=width,
                carton_height=height,
                carton_weight=weight,
                product_verification='Approved',
                min_order_quantity=min_order_quantity,
                price=price,
                discount=0,
                final_price=final_price,
                status='bulk',
                refpro=refpro,
                reuser=reuser,
                product_type_id=product_type_obj.id,
                product_packaging_id=packaging_obj.id,
                product_packaging_value=product_packaging_value,
                shipping_via=shipping_via,
            )

            # -------- MANY TO MANY -------- #
            if list_available_country:
                product.available_country.set(list_available_country)

            if list_tags:
                product.tag.set(list_tags)

            # -------- MODEL -------- #
            model_obj = models.ProductModel.objects.create(
                product_id=product.id,
                model_name=product_name,
                model_name_french=product_name,
                status='Active'
            )

            # -------- VARIANT -------- #
            variant = models.ProductModelVariant.objects.create(
                product_id=product.id,
                model_id=model_obj.id,
                name=product_name,
                name_french=product_name,
                price=final_price,
                variant_verification='Approved',
                status='Active'
            )

            # -------- VENDOR PRICE -------- #
            models.VendorProductPrice.objects.create(
                product_id=product.id,
                model_id=model_obj.id,
                vendor_id=vendor_obj.id,
                variant_id=variant.id,
                price=price,
                status='Active'
            )

            # -------- UPDATE FINAL PRICE -------- #
            product.price = final_price
            product.save()

        return HttpResponse(
            JSONRenderer().render({
                "statuscode": 200,
                "status": "Success",
                "message": "Data submitted successfully"
            }),
            content_type='application/json'
        )

    except Exception as e:
        return HttpResponse(
            JSONRenderer().render({
                "status": "Error",
                "message": str(e)
            }),
            content_type='application/json'
        )


############## S3 File Store end here ##############
#####          DO NOT DELETE THIS CODE           #####
######################################################



@csrf_exempt
def update_vector(request):
    if request.method == "POST":
        all_products = models.ProductDetail.objects.all().order_by('id')

        for product in all_products:
            print("product------>",product)
            get_product_image_1 = product.product_image_1
            get_product_image_2 = product.product_image_2
            get_product_image_3 = product.product_image_3
            get_product_image_4 = product.product_image_4
            get_product_image_5 = product.product_image_5
            get_product_image_6 = product.product_image_6
            get_product_image_7 = product.product_image_7
            get_product_image_8 = product.product_image_8

            domain_path = ""


            if get_product_image_1 != None and get_product_image_1 != 'null' and get_product_image_1 != '':
                feat = extract_features_from_url(domain_path+get_product_image_1.url)
                try:
                    if feat.size > 0:
                        product.product_image_1_vector = feat.tolist()
                        product.save()
                except:
                    pass
            
            if get_product_image_2 != None and get_product_image_2 != 'null' and get_product_image_2 != '':
                feat = extract_features_from_url(domain_path+get_product_image_2.url)
                try:
                    if feat.size > 0:
                        product.product_image_2_vector = feat.tolist()
                        product.save()
                except:
                    pass

            if get_product_image_3 != None and get_product_image_3 != 'null' and get_product_image_3 != '':
                feat = extract_features_from_url(domain_path+get_product_image_3.url)
                try:
                    if feat.size > 0:
                        product.product_image_3_vector = feat.tolist()
                        product.save()
                except:
                    pass

            if get_product_image_4 != None and get_product_image_4 != 'null' and get_product_image_4 != '':
                feat = extract_features_from_url(domain_path+get_product_image_4.url)
                try:
                    if feat.size > 0:
                        product.product_image_4_vector = feat.tolist()
                        product.save()
                
                except:
                    pass


            if get_product_image_5 != None and get_product_image_5 != 'null' and get_product_image_5 != '':
                feat = extract_features_from_url(domain_path+get_product_image_5.url)
                try:
                    if feat.size > 0:
                        product.product_image_5_vector = feat.tolist()
                        product.save()
                except:
                    pass

            if get_product_image_6 != None and get_product_image_6 != 'null' and get_product_image_6 != '':
                feat = extract_features_from_url(domain_path+get_product_image_6.url)
                try:
                    if feat.size > 0:
                        product.product_image_6_vector = feat.tolist()
                        product.save()
                except:
                    pass

            if get_product_image_7 != None and get_product_image_7 != 'null' and get_product_image_7 != '':
                feat = extract_features_from_url(domain_path+get_product_image_7.url)
                try:
                    if feat.size > 0:
                        product.product_image_7_vector = feat.tolist()
                        product.save()
                except:
                    pass

            if get_product_image_8 != None and get_product_image_8 != 'null' and get_product_image_8 != '':
                feat = extract_features_from_url(domain_path+get_product_image_8.url)
                try:
                    if feat.size > 0:
                        product.product_image_8_vector = feat.tolist()
                        product.save()
                except:
                    pass

        res={
            'message':'Product Vector Updated for all image'
        }
        return HttpResponse(JSONRenderer().render(res), content_type="application/json", status=200)


@csrf_exempt
def all_bulk_product_list(request):
    python_data = JSONParser().parse(io.BytesIO(request.body))
    
    print("python_data--all_product_list--->", python_data)

    today = timezone.now()

    page_number = int(python_data.get('page_number', 1))
    row_size = int(python_data.get('row_data', 10))
    last_row = row_size * page_number
    first_row = last_row - row_size

    category_name = python_data.get('category_name',None)
    search_key = python_data.get('search_key',None)
    subcategory_name = python_data.get('subcategory_name',None)

    filter_condition = Q()

    if category_name not in [None,'','null']:
        filter_condition &= Q(category__category = category_name)
    if subcategory_name not in [None,'','null']:
        filter_condition &= Q(subcategory__subcategory = subcategory_name)
    if search_key not in [None,'','null']:
        filter_condition &= Q(product_name__icontains = search_key)|Q(refpro__icontains = search_key)
    
    total_records = models.ProductDetail.objects.filter(filter_condition, status= 'bulk').count()
    print(total_records, 'total_records')
    list_data = []
    all_product = models.ProductDetail.objects.filter(filter_condition, status= 'bulk').order_by('-id')[first_row:last_row]
    for product in all_product:
        product = models.ProductDetail.objects.get(id = product.id)
        price_data = models.ProductModelVariant.objects.filter(product=product).aggregate(
            max_price=Max('price'),
            min_price=Min('price')
        )

        max_price = price_data['max_price']
        min_price = price_data['min_price']
        serializer = ProductDetailSerializer(product).data
        serializer.update({'max_price': max_price, 'min_price':min_price})
        list_data.append(serializer)

    
    res = {
        'data': list_data ,
        'total_records':total_records,
        'current_page':page_number,
        'total_pages': int(np.ceil(total_records/row_size))
        }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)


@csrf_exempt
def product_verification_update(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)
    id = python_data.get('id')

    request_form = models.ProductDetail.objects.get(id = id)
    request_form.product_verification = python_data.get('product_verification', request_form.product_verification)
    request_form.save()

    res = {
        'message':'Product verification update successfully.'        
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)

@csrf_exempt
def pending_product_verification_list_admin(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)
    
    page_number = int(python_data.get('page_number', 1))
    row_size = int(python_data.get('row_data', 10))
    last_row = row_size * page_number
    first_row = last_row - row_size

    total_pending_count = models.ProductDetail.objects.filter(product_verification = 'Pending').count()
    all_product = models.ProductDetail.objects.filter(product_verification = 'Pending')[first_row:last_row]
    serializer = ProductDetailSerializer(all_product, many=True).data

    res = {
        'data': serializer,
        'total_pending_product_count':total_pending_count
    }
    json_data = JSONRenderer().render(res)
    return HttpResponse(json_data, content_type= 'application/json', status=200)


@csrf_exempt
def subcategory_wise_product_list(request):
    if request.method == 'POST':
        python_data = JSONParser().parse(io.BytesIO(request.body))
        # print("python_data--subcategory_wise_product_list-->",python_data)

        currencyCode = python_data.get('currencyCode',None)
        countryCallingCode = python_data.get('countryCallingCode',None)
        country = python_data.get('countryName','Egypt')
        ip_address = python_data.get('ip_address')
        user_id = python_data.get('user_id')

        try:
            if currencyCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(currency_code = currencyCode).first().country_name
            elif countryCallingCode not in [None,'','null']:
                country = models.CountryWithCurrency.objects.filter(country_calling_code = countryCallingCode).first().country_name
                
            elif user_id:
                user = models.CustomerDetail.objects.get(id = user_id)
                countryCode = user.countryCode

                if models.CountryWithCurrency.objects.filter(country_calling_code = countryCode).exists():
                    country = models.CountryWithCurrency.objects.get(country_calling_code = countryCode).country_name
                else:
                    if country not in [None,'','null']:
                        country = country
                    else:
                        try:
                            response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                            response = response.json()
                            country = response.get("country_name")
                            # print("response====>",response)

                            if country in [None,'','null']:
                                country = "Egypt"

                        except Exception as e:
                            # print("Errroooorrr----->",e)
                            country = "Egypt"
                        

            else:
                if country in [None,'','null']:
                    try:
                        response = requests.get(f'https://ipapi.co/{ip_address}/json/')
                        response = response.json()
                        country = response.get("country_name")
                        # print("response====>",response)

                        if country in [None,'','null']:
                            country = "Egypt"

                    except Exception as e:
                        # print("Errroooorrr----->",e)
                        country = "Egypt"
                else:
                    country = country
        except Exception as e:
            # print("Error-=-=-=--->",e)
            country = 'Egypt'

        subcategory = python_data.get('subcategory')
        subcategory_list = models.ProductDetail.objects.filter(subcategory = subcategory, status = 'Active')
        subcategory_serializer = ProductSearchSerializer(subcategory_list, many=True, context = {'country':country}).data
        res={
            'data':subcategory_serializer
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status=200)


@csrf_exempt
def product_list_tag_result(request):
    if request.method == "POST":
        python_data = JSONParser().parse(io.BytesIO(request.body))
        print("python_data---->", python_data)

        tag = python_data.get('tag')
        country = python_data.get('country')
        search_by = python_data.get('search_by')
        search = python_data.get('search')
        currencyCode = python_data.get('currencyCode')

        try:
            if currencyCode not in [None, '', 'null']:
                country_obj = models.CountryWithCurrency.objects.filter(currency_code=currencyCode).first()
                country_name = country_obj.country_name if country_obj else "Egypt"
            else:
                country_name = "Senegal"
        except Exception:
            country_name = "Egypt"

        tag_result_serializer = []
        base_country_query = Q()

        if country not in [None, '', 'null'] and isinstance(country, int):
            base_country_query &= Q(available_country__id=country)

            # =======================================================
            # TAG / SEARCH BY FILTER (FIXED & LINKED TO BASE QUERY)
            # =======================================================
            if tag not in [None, '', 'null']:
                search_query = Q()  # Initialized to prevent NameError
                
                if search_by == "Tag":
                    search_query &= Q(tag__id=tag)
                elif search_by == "product_english":
                    product_obj = models.ProductDetail.objects.filter(id=tag).first()
                    print(product_obj, 'product_obj')
                    if product_obj:
                        search_query &= Q(product_name__icontains=product_obj.product_name)
                elif search_by == "product_french":
                    product_obj = models.ProductDetail.objects.filter(id=tag).first()
                    if product_obj and product_obj.product_name_french:
                        search_query &= Q(product_name_french__icontains=product_obj.product_name_french)
                
                # Append the search_by filters directly to your country query
                base_country_query &= search_query

            # =====================================
            # SEARCH CLEANING
            # =====================================
            search_clean = ""
            if search not in [None, '', 'null']:
                search = str(search).strip()
                search_clean = search.replace(" ", "").replace("-", "").replace("_", "").lower()

            # =====================================
            # BASE QUERYSET
            # =====================================
            products = models.ProductDetail.objects.filter(base_country_query, status='Active')

            # =====================================
            # STRICT TEXT & FUZZY SEARCH LOGIC
            # =====================================
            if search_clean:
                normalized_search = re.sub(r'(.)\1+', r'\1', search.lower())
                search_words = normalized_search.split()

                # 1. Build a strict keyword match (ALL words must match somewhere)
                strict_text_query = Q()
                for word in search_words:
                    if len(word) <= 2 and not word.isdigit():
                        continue  # Skip junk words like "du", "le"
                    
                    # Each individual word must hit name, french name, OR refpro
                    word_match = (
                        Q(product_name__unaccent__icontains=word) |
                        Q(product_name_french__unaccent__icontains=word) |
                        Q(refpro__unaccent__icontains=word)
                    )
                    
                    if not strict_text_query:
                        strict_text_query = word_match
                    else:
                        strict_text_query &= word_match  # AND condition forces all words to match

                # 2. Add Trigram scoring metrics for ranking relevance
                products = products.annotate(
                    similarity_name=TrigramSimilarity('product_name', normalized_search),
                    similarity_name_word=TrigramWordSimilarity(normalized_search, 'product_name'),
                    similarity_name_french=TrigramSimilarity('product_name_french', normalized_search),
                    similarity_name_french_word=TrigramWordSimilarity(normalized_search, 'product_name_french'),
                    similarity_refpro=TrigramSimilarity('refpro', normalized_search),
                    similarity_refpro_word=TrigramWordSimilarity(normalized_search, 'refpro')
                ).annotate(
                    best_similarity=Greatest(
                        'similarity_name', 'similarity_name_word',
                        'similarity_name_french', 'similarity_name_french_word',
                        'similarity_refpro', 'similarity_refpro_word'
                    )
                )

                # Execute primary strict search first
                tag_result = products.filter(strict_text_query).distinct().order_by('-best_similarity', 'product_name')

                # If strict filtering returns nothing, fall back to high-confidence fuzzy match
                if not tag_result.exists():
                    tag_result = products.filter(best_similarity__gt=0.40).distinct().order_by('-best_similarity', 'product_name')
            else:
                tag_result = products.distinct().order_by('product_name')

            # Serialize primary results
            tag_result_serializer = ProductSearchSerializer(
                tag_result, many=True, context={'country': country_name}
            ).data

            # =====================================
            # FALLBACK SEARCH (IF NO RESULTS FOUND)
            # =====================================
            if not tag_result_serializer and tag not in [None, '', 'null']:
                # Setup cleaner string manipulations on the fields for fallback matching
                products_clean_fields = models.ProductDetail.objects.filter(
                    available_country__id=country, status='Active'
                ).annotate(
                    clean_product_name=Lower(Replace(Replace(Replace('product_name', Value(' '), Value('')), Value('-'), Value('')), Value('_'), Value(''))),
                    clean_product_name_french=Lower(Replace(Replace(Replace('product_name_french', Value(' '), Value('')), Value('-'), Value('')), Value('_'), Value(''))),
                    clean_refpro=Lower(Replace(Replace(Replace('refpro', Value(' '), Value('')), Value('-'), Value('')), Value('_'), Value('')))
                )

                try:
                    tag_name = models.ProductTag.objects.get(id=tag).tag
                    fallback_search = re.sub(r'(.)\1+', r'\1', tag_name.lower()).replace(" ", "").replace("-", "").replace("_", "")

                    without_tag_result = products_clean_fields.filter(
                        Q(clean_product_name__icontains=fallback_search) |
                        Q(clean_product_name_french__icontains=fallback_search) |
                        Q(clean_refpro__icontains=fallback_search)
                    ).distinct()

                    tag_result_serializer = ProductSearchSerializer(
                        without_tag_result, many=True, context={'country': country_name}
                    ).data

                except Exception:
                    if search_clean:
                        without_tag_result = products_clean_fields.filter(
                            Q(clean_product_name__icontains=search_clean) |
                            Q(clean_product_name_french__icontains=search_clean) |
                            Q(clean_refpro__icontains=search_clean)
                        ).distinct()

                        tag_result_serializer = ProductSearchSerializer(
                            without_tag_result, many=True, context={'country': country_name}
                        ).data

        res = {"data": tag_result_serializer}
        return HttpResponse(JSONRenderer().render(res), content_type="application/json", status=200)


@csrf_exempt
def model_update(request):
    if request.method == "POST":
        python_data={}
        for i,j in request.FILES.items():
            python_data.update({i:j})
        
        for i,j in request.POST.items():
            python_data.update({i:j})

        id = python_data.get('id')
        model_name = python_data.get('model_name')
        model_name_french = python_data.get('model_name_french')
        model_image = python_data.get('model_image',None)

        if models.ProductModel.objects.filter(id = id):
            get_model = models.ProductModel.objects.get(id = id)
            get_model.model_name = model_name
            get_model.model_name_french = model_name_french
            
            if not isinstance(model_image, str) and model_image not in [None,'','null']:
                get_model.model_image = model_image

            get_model.save()
            res = {
                'message':'Model Updated Successfully'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=200)

        else:
            res = {
                'message':'Enter Valid id.'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=406)


@csrf_exempt
def variant_update(request):
    if request.method == "POST":
        python_data={}
        for i,j in request.FILES.items():
            python_data.update({i:j})
        
        for i,j in request.POST.items():
            python_data.update({i:j})

        id = python_data.get('id')
        name = python_data.get('name')
        name_french = python_data.get('variant_name_french')
        image = python_data.get('image')
        price = python_data.get('price')

        if models.ProductModelVariant.objects.filter(id = id):
            get_variant = models.ProductModelVariant.objects.get(id = id)
            get_variant.name = name
            get_variant.name_french = name_french
            get_variant.price = price
            if not isinstance(image, str) and image not in [None,'','null']:
                get_variant.image = image

            get_variant.save()
            print(get_variant.product)
            try:
                list_price = list(models.ProductModelVariant.objects.filter(product = get_variant.product.id).values_list('price', flat = True))
                smallest_price = min(list_price)
                # print(smallest_price, 'smallest_pricesmallest_pricesmallest_price')
                get_product = models.ProductDetail.objects.get(id = get_variant.product.id)
                get_product.price = smallest_price
                get_product.save()
            except:
                pass

            res = {
                'message':'Variant Updated Successfully'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=200)

        else:
            res = {
                'message':'Enter Valid id.'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=406)


@csrf_exempt
def vendor_variant_update(request):
    if request.method == "POST":
        python_data={}
        for i,j in request.FILES.items():
            python_data.update({i:j})
        
        for i,j in request.POST.items():
            python_data.update({i:j})

        id = python_data.get('id')
        price = python_data.get('price')
        

        if models.VendorProductPrice.objects.filter(id = id):
            get_vendor_variant = models.VendorProductPrice.objects.get(id = id)
            get_vendor_variant.price = price

            get_vendor_variant.save()
            res = {
                'message':'Variant Price Updated Successfully'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=200)

        else:
            res = {
                'message':'Enter Valid id.'        
            }
            return HttpResponse(JSONRenderer().render(res), content_type= 'application/json', status=406)



@csrf_exempt
def all_product_model_variant_list_admin(request):
    if request.method == "POST":
        product_model_variant_list = models.ProductDetail.objects.all().order_by('-id')
        product_model_variant_list_serializer = AdminProductDetailSerializer(product_model_variant_list, many=True).data

        res={
            'data':product_model_variant_list_serializer
        }
        return HttpResponse(JSONRenderer().render(res), content_type='application/json', status=200)



@csrf_exempt
def get_product_data_update(request):
    if request.method =="POST":
        
        product_list = models.ProductDetail.objects.all().order_by('id')[:200]
        for product in product_list:
            print(product.id, 'product.id')
            product = models.ProductDetail.objects.get(id = product.id)
            get_all_variant_price = models.ProductModelVariant.objects.filter(product = product.id).values_list('price', flat=True)
            min_price = min(get_all_variant_price)
            # print(min_price, product.product_name)
            product.price = min_price
            product.save()


        res={
            'message':'Product Status update Succesfully.' 
        }
        return HttpResponse(JSONRenderer().render(res), content_type = 'application/json', status = 200)


@csrf_exempt
def all_product_tag_list(request):
    json_data = request.body
    stream = io.BytesIO(json_data)
    python_data = JSONParser().parse(stream)

    page_number = int(python_data.get('page_number', 1))
    row_size = int(python_data.get('row_data', 10))

    cache_key = f"product_tags_page_{page_number}_size_{row_size}"

    # Check cache first
    cached_data = cache.get(cache_key)
    if cached_data:
        return HttpResponse(
            JSONRenderer().render(cached_data),
            content_type='application/json',
            status=200
        )

    last_row = row_size * page_number
    first_row = last_row - row_size

    total_tags = models.ProductTag.objects.count()

    all_product = models.ProductTag.objects.all().order_by('-id')[first_row:last_row]

    serializer = ProductTagSerializer(all_product, many=True).data

    res = {
        'data': serializer,
        'total_tags': total_tags
    }

    # Save response in Redis for 5 minutes
    cache.set(cache_key, res, timeout=300)

    json_data = JSONRenderer().render(res)

    return HttpResponse(json_data, content_type='application/json', status=200)